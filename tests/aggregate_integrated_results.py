# -*- coding: utf-8 -*-
"""5명이 회신한 gsnd_consistency_sheet{N}_*.xlsx 들을 한 파일로 취합.

각 결과 파일의 `diff`(단계별 흔들림)·`변동상세`(흔들린 값)를 모아:
- summary    : 시트별 질문수·흔들린질문수·단계별 흔들림 건수 한눈 요약(+합계)
- diff_all   : 5개 diff 를 시트 컬럼 붙여 이어붙임(흔들린 셀 노랑 강조)
- 변동상세_all: 5개 변동상세를 이어붙임(회차 셀 강조)

같은 시트번호 파일이 여러 개면 파일명 최신(타임스탬프) 것만 사용.

사용:
  python tests/aggregate_integrated_results.py                 # qa_test/ 에서 수집
  python tests/aggregate_integrated_results.py --dir 결과폴더   # 지정 폴더에서 수집
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
FILE_RE = re.compile(r"gsnd_consistency_sheet(\d+)_", re.I)
_FLAP_FILL = PatternFill(start_color="FFF2A2", end_color="FFF2A2", fill_type="solid")
_HDR_FONT = Font(bold=True)
SHEET_LABEL = {1: "DQP싱글턴", 2: "DQP멀티턴", 3: "경남셋1", 4: "경남셋2", 5: "DQP추가"}


def find_files(d: Path) -> dict[int, Path]:
    """시트번호 → 최신 파일. 같은 번호 중복은 타임스탬프 최신만."""
    cand: dict[int, list[Path]] = {}
    for p in sorted(d.glob("gsnd_consistency_sheet*_*.xlsx")):
        m = FILE_RE.search(p.name)
        if m:
            cand.setdefault(int(m.group(1)), []).append(p)
    picked: dict[int, Path] = {}
    for n, ps in sorted(cand.items()):
        ps.sort()
        picked[n] = ps[-1]
        for extra in ps[:-1]:
            print(f"  [중복] 시트{n}: 최신만 사용 {ps[-1].name} (무시 {extra.name})")
    return picked


def _read_sheet(path: Path, name: str):
    """(헤더, 데이터행들) — 시트 없으면 (None, [])."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        wb.close()
        return None, []
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    if not rows:
        return None, []
    header = list(rows[0])
    data = [list(r) for r in rows[1:] if r and r[0] not in (None, "")]
    return header, data


def _is_flap(cell) -> bool:
    return "⚠" in str(cell) if cell is not None else False


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(ROOT / "qa_test"), help="결과 xlsx 수집 폴더(기본 qa_test/)")
    ap.add_argument("--out", default=None, help="출력 경로(기본 <dir>/gsnd_consistency_취합_<ts>.xlsx)")
    args = ap.parse_args()

    src_dir = Path(args.dir)
    files = find_files(src_dir)
    if not files:
        raise SystemExit(f"[error] gsnd_consistency_sheet*_*.xlsx 없음: {src_dir}")
    missing = [n for n in range(1, 6) if n not in files]
    print(f"[취합] 수집 {len(files)}개: 시트 {sorted(files)}"
          + (f"  ⚠미도착 시트 {missing}" if missing else "  (1~5 모두 도착)"))

    out = openpyxl.Workbook()

    # ---- diff_all (이어붙임) ----
    wda = out.active
    wda.title = "diff_all"
    stages: list[str] = []
    summary_rows: list[list] = []      # (시트, 라벨, 파일, 질문수, 흔들린질문, 안정질문, *단계별흔들림)
    header_written = False

    for n in sorted(files):
        header, data = _read_sheet(files[n], "diff")
        if header is None:
            print(f"  [경고] 시트{n}: diff 없음 → 건너뜀 ({files[n].name})")
            continue
        cur_stages = [str(h) for h in header[2:]]
        if not stages:
            stages = cur_stages
        if not header_written:
            wda.append(["시트", "라벨"] + [str(h) for h in header])  # 질문/회차수/단계들
            for c in wda[1]:
                c.font = _HDR_FONT
            header_written = True

        flapped_q = 0
        per_stage = [0] * len(cur_stages)
        for r in data:
            row_stages = r[2:]
            flaps = [_is_flap(c) for c in row_stages]
            if any(flaps):
                flapped_q += 1
            for i, f in enumerate(flaps):
                if i < len(per_stage) and f:
                    per_stage[i] += 1
            wda.append([n, SHEET_LABEL.get(n, "")] + r)
            xr = wda.max_row
            for i, f in enumerate(flaps):
                if f:
                    wda.cell(row=xr, column=5 + i).fill = _FLAP_FILL  # 시트,라벨,질문,회차수 다음
        nq = len(data)
        summary_rows.append([n, SHEET_LABEL.get(n, ""), files[n].name, nq,
                             flapped_q, nq - flapped_q] + per_stage)

    # ---- summary ----
    ws = out.create_sheet("summary", 0)
    ws.append(["시트", "라벨", "파일", "질문수", "흔들린질문", "안정질문"] + stages)
    for c in ws[1]:
        c.font = _HDR_FONT
    tot = [0] * (3 + len(stages))   # 질문수,흔들린,안정 + 단계들
    for row in summary_rows:
        ws.append(row)
        nums = row[3:]
        for i, v in enumerate(nums):
            tot[i] += v if isinstance(v, int) else 0
    ws.append(["합계", "", ""] + tot)
    for c in ws[ws.max_row]:
        c.font = _HDR_FONT

    # ---- 변동상세_all ----
    wf = out.create_sheet("변동상세_all")
    detail_header_written = False
    for n in sorted(files):
        header, data = _read_sheet(files[n], "변동상세")
        if header is None:
            continue
        if not detail_header_written:
            wf.append(["시트", "라벨"] + [str(h) for h in header])
            for c in wf[1]:
                c.font = _HDR_FONT
            detail_header_written = True
        for r in data:
            wf.append([n, SHEET_LABEL.get(n, "")] + r)
            xr = wf.max_row
            # 원본 r[j] → 출력열 = 시트·라벨(2) + (j+1). R1=r[3]부터 강조
            for j in range(3, len(r)):
                wf.cell(row=xr, column=j + 3).fill = _FLAP_FILL

    # 열너비 정리
    for sh in (ws, wda, wf):
        for col in range(1, sh.max_column + 1):
            sh.column_dimensions[get_column_letter(col)].width = 18

    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = Path(args.out) if args.out else src_dir / f"gsnd_consistency_취합_{ts}.xlsx"
    out.save(out_path)

    # 콘솔 요약
    print("\n시트 라벨        질문수  흔들린  안정")
    for row in summary_rows:
        print(f"  {row[0]} {row[1]:<10} {row[3]:>5} {row[4]:>6} {row[5]:>6}")
    print(f"\n[done] → {out_path}")


if __name__ == "__main__":
    main()
