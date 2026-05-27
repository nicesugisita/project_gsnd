"""사업명 set 골든셋 템플릿 생성 — 데이터셋 final_answer(과거 스냅샷)에서 사업명을 초안으로 추출.

사람이 '기대사업명(교정)' 칸을 빼고/더해 gold 를 확정한 뒤, score_against_gold.py 로 채점한다.
스냅샷은 정답(gold)이 아니라 '초안 시드'일 뿐 — 반드시 사람 검토를 거친다.

사용:
    python tests/build_gold_template.py [데이터셋_xlsx] [출력_xlsx]
기본: qa_test/260513_gsnd_total_dataset_v0.2.xlsx → qa_test/gold_services_v0.xlsx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_checklist import _service_names  # noqa: E402

DEFAULT_DATASET = ROOT / "qa_test" / "260513_gsnd_total_dataset_v0.2.xlsx"
DEFAULT_OUT = ROOT / "qa_test" / "gold_services_v0.xlsx"
_MT = re.compile(r"^\[(\d+)-(\d+)\]")

HDR = PatternFill("solid", fgColor="DDEBF7")
DRAFT_FILL = PatternFill("solid", fgColor="F2F2F2")  # 초안=읽기참고(회색)
EDIT_FILL = PatternFill("solid", fgColor="FFF7E0")   # 교정칸=입력대상(연노랑)
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

# (제목, 너비)
_COLS = [
    ("no", 6), ("question", 40), ("clarified", 9), ("follow_up", 10),
    ("스냅샷초안(사업명·참고)", 34), ("기대사업명(교정·줄바꿈구분)", 34),
    ("검토", 10), ("비고", 28),
]


def main() -> None:
    dataset = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DATASET
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    if not dataset.exists():
        print(f"[error] 데이터셋 없음: {dataset}"); sys.exit(1)

    wb_in = openpyxl.load_workbook(dataset, read_only=True, data_only=True)
    ws_in = wb_in["gsnd_total"]

    wb = Workbook()
    ws = wb.active
    ws.title = "gold"
    ws.append([c[0] for c in _COLS])
    for j, c in enumerate(_COLS, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = HDR_FONT; cell.fill = HDR; cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(j)].width = c[1]

    n_total = n_card = 0
    for i, r in enumerate(ws_in.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None or r[1] is None:
            continue
        q = str(r[1]).strip()
        if not q or _MT.match(q):
            continue  # 멀티턴 연속행 제외 (러너와 동일)
        try:
            no = int(r[0])
        except (TypeError, ValueError):
            continue
        clarified = str(r[2]).strip() if len(r) > 2 and r[2] else ""
        follow_up = str(r[3]).strip() if len(r) > 3 and r[3] else ""
        ans = str(r[4]) if len(r) > 4 and r[4] is not None else ""
        draft = _service_names(ans)
        n_total += 1
        if draft:
            n_card += 1
        draft_str = "\n".join(draft)
        ws.append([no, q, clarified, follow_up, draft_str, draft_str, "", ""])
        rr = ws.max_row
        ws.cell(row=rr, column=5).fill = DRAFT_FILL
        ws.cell(row=rr, column=6).fill = EDIT_FILL
        for cc in (2, 5, 6, 8):
            ws.cell(row=rr, column=cc).alignment = WRAP

    last = ws.max_row
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLS))}{last}"
    dv = DataValidation(type="list", formula1='"미검토,검토완료,제외"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")

    # 안내 시트
    g = wb.create_sheet("사용법", 0)
    g.column_dimensions["A"].width = 100
    for line in [
        "■ 사업명 set 골든셋 — 작성 안내",
        "",
        "1) '기대사업명(교정)' 칸이 채점 기준입니다. '스냅샷초안'은 과거 출력에서 뽑은 참고용(정답 아님).",
        "2) 초안을 검토해 ▶틀린/무관한 사업명 삭제, ▶빠진 사업명 추가. 한 줄에 사업명 하나(줄바꿈 구분).",
        "3) 연도·지역 접두어(예 '2026년 창원시')는 적어도 무방 — 채점 시 자동 정규화됩니다.",
        "4) '검토' 칸: 미검토/검토완료/제외. '제외'는 채점에서 빼고 싶은 문항(예: 추천형 아님).",
        "5) 비추천형(단일조회·비교 등)으로 기대 사업명이 없으면 '기대사업명'을 비우고 검토=제외.",
        "",
        f"■ 채점: python tests/score_against_gold.py {DEFAULT_OUT.name} <consistency.jsonl>",
        "   → 문항별 Precision(이상한거 안섞임)·Recall(관련 다나옴)·F1 + 요약",
        "",
        "※ 스냅샷은 정답이 아닙니다. 초안을 그대로 두면 '과거출력과의 일치'를 재는 셈이니 반드시 교정하세요.",
    ]:
        g.append([line])
    g["A1"].font = Font(bold=True, size=13)

    wb.save(out)
    print(f"[done] gold 템플릿 → {out}")
    print(f"  문항 {n_total} (사업명 초안 있음 {n_card} / 초안 없음 {n_total - n_card})")
    print("  '기대사업명' 칸을 교정한 뒤 score_against_gold.py 로 채점하세요.")


if __name__ == "__main__":
    main()
