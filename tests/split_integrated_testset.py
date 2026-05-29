# -*- coding: utf-8 -*-
"""260528 통합테스트셋 → 시트별 gsnd_total 스키마 xlsx 5개로 분리.

run_gsnd_consistency.py 가 `--src` 로 그대로 소비하는 형식
(시트명 `gsnd_total` / 컬럼 no·question·clarified·follow_up)으로 변환한다.

- 단일턴 시트(1·3·5): question 컬럼만 뽑아 평면 변환.
- 멀티턴 시트(2): 이미 `[N-M]` 프리픽스가 있어 그대로 통과(clarified/follow_up 보존).
- 멀티턴 시트(4): 번호 `N-M` 를 읽어 turn>1 행에 `[N-M]` 프리픽스를 부여.
  → --multiturn-only replay 시 대화 단위로 묶인다(turn1 no == base 번호).

사용:  python tests/split_integrated_testset.py
출력:  tests/integrated_split/sheet{1..5}_*.xlsx   (git 추적 폴더 — 수신자 전달용)
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "qa_test" / "260528_통합테스트셋_DQ-경남_v1.3.xlsx"
OUT_DIR = ROOT / "tests" / "integrated_split"   # git 추적 폴더(qa_test/* 는 gitignore) — 수신자 전달용
DEFAULT_SIGUN = "창원"   # follow_up 미작성 시 시/군 되묻기에 사용할 기본 답변

_BULLET = re.compile(r"^[\s　•·▪◦・]+")          # 선두 불릿/전각공백만 제거(하이픈·숫자 보존)
_NUM = re.compile(r"^\s*(\d+)\s*-\s*(\d+)")           # "N-M" 번호 파싱
_PLACEHOLDER = re.compile(r"[○◯〇]+\s*시군")          # "○○시군" 자리표시자 → 창원

# 시트별 설정. col_* 는 0-based 컬럼 인덱스.
SHEETS = [
    dict(idx=1, name="(DQP)싱글턴 테스트셋_150건", out="sheet1_dqp_single",
         mode="single", header=1, c_no=0, c_q=2, c_clar=3, c_fu=4),
    dict(idx=2, name="(DQP)멀티턴 테스트셋_50건(건당5턴)", out="sheet2_dqp_multi",
         mode="passthrough", header=3, c_no=0, c_q=2, c_clar=3, c_fu=4),
    dict(idx=3, name="(경남)테스트셋1_100건", out="sheet3_gn_set1",
         mode="single", header=2, c_no=0, c_q=3, c_clar=None, c_fu=None),
    dict(idx=4, name="(경남)테스트셋2_19건(3턴)", out="sheet4_gn_set2",
         mode="numbered", header=1, c_num=1, c_q=2),
    dict(idx=5, name="(DQP)추가생성데이터_510건", out="sheet5_dqp_extra",
         mode="single", header=1, c_no=0, c_q=1, c_clar=None, c_fu=None),
]


def _get(row, i):
    if i is None or i >= len(row):
        return None
    return row[i]


def _clean(v) -> str:
    if v is None:
        return ""
    t = str(v).replace("　", " ").strip()
    t = _BULLET.sub("", t).strip()
    return _PLACEHOLDER.sub("창원", t)


def _norm_keep_prefix(v) -> str:
    """[N-M] 프리픽스를 보존하는 가벼운 정규화(passthrough 용)."""
    if v is None:
        return ""
    return _PLACEHOLDER.sub("창원", str(v).replace("　", " ").strip())


def _yes(v) -> bool:
    return str(v).strip().upper() in ("Y", "YES", "TRUE", "1") if v is not None else False


def _as_int(v, default):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def build_rows(ws, cfg) -> list[tuple]:
    """(no, question, clarified, follow_up) 리스트로 변환."""
    data = list(ws.iter_rows(min_row=cfg["header"] + 1, values_only=True))
    out: list[tuple] = []
    seq = 0
    for row in data:
        mode = cfg["mode"]
        if mode == "numbered":
            num = str(_get(row, cfg["c_num"]) or "").strip()
            q = _clean(_get(row, cfg["c_q"]))
            if not q:
                continue
            m = _NUM.match(num)
            if not m:
                continue
            base, turn = int(m.group(1)), int(m.group(2))
            if turn == 1:
                out.append((base, q, "", DEFAULT_SIGUN))       # turn1: no=base, 프리픽스 없음
            else:
                out.append((base * 100 + turn, f"[{base}-{turn}] {q}", "", ""))
        elif mode == "passthrough":
            q = _norm_keep_prefix(_get(row, cfg["c_q"]))        # [N-M] 프리픽스 유지
            if not q:
                continue
            seq += 1
            no = _as_int(_get(row, cfg["c_no"]), seq)
            clar = "Y" if _yes(_get(row, cfg["c_clar"])) else ""
            fu = _clean(_get(row, cfg["c_fu"])) or DEFAULT_SIGUN  # 비면 '창원'으로 되묻기 응답
            out.append((no, q, clar, fu))
        else:  # single
            q = _clean(_get(row, cfg["c_q"]))
            if not q:
                continue
            seq += 1
            no = _as_int(_get(row, cfg["c_no"]), seq)
            clar = "Y" if _yes(_get(row, cfg.get("c_clar"))) else ""
            fu = (_clean(_get(row, cfg.get("c_fu"))) if cfg.get("c_fu") is not None else "") or DEFAULT_SIGUN
            out.append((no, q, clar, fu))
    return out


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"[error] 입력 파일 없음: {SRC}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    print(f"[split] src={SRC.name}")
    for cfg in SHEETS:
        if cfg["name"] not in wb.sheetnames:
            print(f"  [skip] 시트 없음: {cfg['name']}")
            continue
        rows = build_rows(wb[cfg["name"]], cfg)
        ob = openpyxl.Workbook()
        ws = ob.active
        ws.title = "gsnd_total"
        ws.append(["no", "question", "clarified", "follow_up"])
        for r in rows:
            ws.append(list(r))
        path = OUT_DIR / f"{cfg['out']}.xlsx"
        ob.save(path)
        # 멀티턴이면 대화(base) 수도 표시
        is_mt = cfg["mode"] in ("passthrough", "numbered")
        if is_mt:
            bases = set()
            for no, q, _, _ in rows:
                m = re.match(r"^\[(\d+)-(\d+)\]", q)
                bases.add(int(m.group(1)) if m else no)
            extra = f"  대화={len(bases)}개"
        else:
            extra = ""
        print(f"  시트{cfg['idx']}: {len(rows):>4}행{extra}  → {path.relative_to(ROOT)}")
    wb.close()
    print(f"[done] {OUT_DIR.relative_to(ROOT)} 아래 5개 파일 생성")


if __name__ == "__main__":
    main()
