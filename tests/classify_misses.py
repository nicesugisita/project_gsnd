"""F1 낮은 문항의 '놓친 정답서비스'를 실제 검색문서와 대조해 진짜누락 vs 명칭차이 분류.

매칭(_canon)이 엄격해서 같은 제도인데 표기만 달라 미검출로 잡힌 경우(명칭차이)와,
정말 안 나온 경우(진짜누락)를 명칭 유사도(difflib)로 구분한다. 휴리스틱 — 최종 확인은 사람.

사용:
    python tests/classify_misses.py <quality.xlsx> [out.xlsx] [F1임계(기본0.3)]
"""
from __future__ import annotations

import re
import sys
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_checklist import _canon  # noqa: E402

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
HDR = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

SIM_THR = 0.55  # 이 이상이면 같은 제도(명칭차이)로 추정


def _n2(s: str) -> str:
    """괄호내용 제거 후 canon (명칭차이 비교용)."""
    s = re.sub(r"[\(\[（].*?[\)\]）]", "", str(s or ""))
    return _canon(s)


def _sim(a: str, b: str) -> float:
    na, nb = _n2(a), _n2(b)
    if not na or not nb:
        return 0.0
    if na in nb or nb in na:           # 핵심 포함 관계
        return max(0.85, SequenceMatcher(None, na, nb).ratio())
    return SequenceMatcher(None, na, nb).ratio()


def _best_match(gold_name: str, shown: list):
    best, who = 0.0, ""
    for d in shown:
        r = _sim(gold_name, d)
        if r > best:
            best, who = r, d
    return best, who


def main() -> None:
    if len(sys.argv) < 2:
        print("usage: python tests/classify_misses.py <quality.xlsx> [out.xlsx] [F1임계]")
        sys.exit(1)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem.replace("_quality", "") + "_miss_classified.xlsx")
    thr = float(sys.argv[3]) if len(sys.argv) > 3 else 0.3

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["분석"]
    H = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(H)}

    targets = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[ci["gold有"]] != "Y":
            continue
        f1 = r[ci["F1"]]
        if isinstance(f1, (int, float)) and f1 < thr:
            targets.append(r)
    targets.sort(key=lambda r: r[ci["F1"]])

    owb = Workbook()
    osh = owb.active
    osh.title = "누락분류"
    cols = [("질문", 38), ("F1", 6), ("놓친 정답서비스", 30), ("판정", 12),
            ("유사 검색문서(명칭차이 추정)", 34), ("유사도", 8)]
    osh.append([c[0] for c in cols])
    for j, c in enumerate(cols, 1):
        cell = osh.cell(row=1, column=j)
        cell.font = HDR_FONT; cell.fill = HDR; cell.alignment = WRAP
        osh.column_dimensions[get_column_letter(j)].width = c[1]

    q_verdict = []  # (질문, F1, n_miss, n_namevar, n_real)
    for r in targets:
        q = r[ci["질문"]]
        f1 = r[ci["F1"]]
        missed = [x.strip() for x in str(r[ci["빠짐(gold-only,미검출)"]] or "").split("|") if x.strip()]
        shown = [x.strip() for x in str(r[ci["검색문서(R1)"]] or "").split("|") if x.strip()]
        n_var = n_real = 0
        for g in missed:
            sim, who = _best_match(g, shown)
            if sim >= SIM_THR:
                verdict, n_var = "명칭차이", n_var + 1
            else:
                verdict, who, sim, n_real = "진짜누락", "", 0.0, n_real + 1
            osh.append([q, round(f1, 2), g, verdict, who, round(sim, 2) if sim else ""])
            rr = osh.max_row
            osh.cell(row=rr, column=4).fill = GREEN if verdict == "명칭차이" else RED
            for cc in (1, 3, 5):
                osh.cell(row=rr, column=cc).alignment = WRAP
        q_verdict.append((q, f1, len(missed), n_var, n_real))

    last = osh.max_row
    osh.freeze_panes = "B2"
    osh.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # 요약 (문항별)
    s = owb.create_sheet("문항요약", 0)
    s.column_dimensions["A"].width = 40
    for w, c in zip((40, 6, 8, 10, 10, 14), "ABCDEF"):
        s.column_dimensions[c].width = w
    s.append(["질문", "F1", "놓친수", "명칭차이", "진짜누락", "종합판정"])
    for c in s[1]:
        c.font = HDR_FONT
    tot_var = tot_real = 0
    for q, f1, nm, nv, nr in q_verdict:
        tot_var += nv; tot_real += nr
        if nv and not nr:
            v = "전부 명칭차이(저평가)"
        elif nr and not nv:
            v = "진짜 누락"
        else:
            v = "혼재"
        s.append([q, round(f1, 2), nm, nv, nr, v])
        rr = s.max_row
        if v.startswith("전부 명칭"):
            s.cell(row=rr, column=6).fill = GREEN
        elif v == "진짜 누락":
            s.cell(row=rr, column=6).fill = RED
        s.cell(row=rr, column=1).alignment = WRAP

    print(f"[done] 불량 {len(q_verdict)}문항 분류 → {out}")
    print(f"  놓친 서비스 총 {tot_var + tot_real}건 중: 명칭차이(실제론 유사 노출) {tot_var} / 진짜누락 {tot_real}")
    owb.save(out)


if __name__ == "__main__":
    main()
