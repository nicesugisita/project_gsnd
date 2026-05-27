"""사업명 set 골든셋 '후보' 생성 — 사람 1패스 확정용 (현재 답변 + 과거 스냅샷).

진짜 정확도(P/R)는 현재 출력에서 파생한 gold로는 못 만든다(편향). 그래서:
- '기대사업명' 프리필 = 공통(두출처 일치) + 스냅샷only  → **현재와 독립적인 베이스**
- '현재추가후보' = 현재only 사업명 (정밀도를 좌우하는 항목) → 사람이 "이 질문에 맞나" 판정,
  채택하면 기대사업명 칸으로 올림. 문서 미근거(환각의심)는 [근거X]로 자동 표시.
이렇게 하면 사람 작업이 "현재only 항목 검토"로 최소화되고, 확정 후 score_against_gold 로 진짜 P/R.

사용:
    python tests/build_gold_candidate.py <consistency_jsonl> [out_xlsx] [dataset_xlsx]
"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_checklist import _service_names, _canon, _docs_blob, _name_grounded  # noqa: E402

DEFAULT_DATASET = ROOT / "qa_test" / "260513_gsnd_total_dataset_v0.2.xlsx"
DEFAULT_OUT = ROOT / "qa_test" / "gold_services_candidate_v0.xlsx"
_MT = re.compile(r"^\[(\d+)-(\d+)\]")

HDR = PatternFill("solid", fgColor="DDEBF7")
EDIT_FILL = PatternFill("solid", fgColor="FFF7E0")   # 기대사업명(채점기준)
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")  # 현재추가후보(검토)
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

_COLS = [
    ("no", 6), ("question", 36),
    ("기대사업명(교정·줄바꿈구분)", 34),
    ("현재추가후보(검토→채택시 위칸으로)", 40),
    ("출처요약", 22), ("검토", 10), ("비고", 20),
]


def _load_snapshot(path: Path) -> Tuple[Dict[int, List[str]], Dict[int, str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["gsnd_total"]
    names: Dict[int, List[str]] = {}
    quest: Dict[int, str] = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None or r[1] is None:
            continue
        q = str(r[1]).strip()
        if not q or _MT.match(q):
            continue
        try:
            no = int(r[0])
        except (TypeError, ValueError):
            continue
        quest[no] = q
        ans = str(r[4]) if len(r) > 4 and r[4] is not None else ""
        names[no] = _service_names(ans)
    return names, quest


def _load_current(path: Path) -> Tuple[Dict[int, List[str]], Dict[int, str], Dict[int, str]]:
    """no → 전 회차 합집합 사업명, 질문, 검색문서 blob(근거 대조용)."""
    acc: Dict[int, List[str]] = defaultdict(list)
    quest: Dict[int, str] = {}
    seen: Dict[int, set] = defaultdict(set)
    docs: Dict[int, list] = defaultdict(list)
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            no = rec.get("_gsnd_no")
            if not isinstance(no, int):
                continue
            quest.setdefault(no, rec.get("user_question", ""))
            for nm in _service_names(rec.get("response_text") or ""):
                c = _canon(nm)
                if c and c not in seen[no]:
                    seen[no].add(c)
                    acc[no].append(nm)
            docs[no].extend(rec.get("top_docs") or [])
    blobs = {no: _docs_blob(docs[no]) for no in docs}
    return dict(acc), quest, blobs


def main() -> None:
    if len(sys.argv) < 2:
        print("usage: python tests/build_gold_candidate.py <consistency_jsonl> [out_xlsx] [dataset_xlsx]")
        sys.exit(1)
    jsonl = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    dataset = Path(sys.argv[3]) if len(sys.argv) > 3 else DEFAULT_DATASET
    if not jsonl.exists():
        print(f"[error] jsonl 없음: {jsonl}"); sys.exit(1)

    snap_names, snap_q = _load_snapshot(dataset)
    cur_names, cur_q, cur_blob = _load_current(jsonl)
    all_no = sorted(set(snap_names) | set(cur_names))

    wb = Workbook()
    ws = wb.active
    ws.title = "gold"
    ws.append([c[0] for c in _COLS])
    for j, c in enumerate(_COLS, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = HDR_FONT; cell.fill = HDR; cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(j)].width = c[1]

    n_rows = n_review = n_ungrounded = 0
    for no in all_no:
        q = cur_q.get(no) or snap_q.get(no) or ""
        s_canon = {_canon(x): x for x in snap_names.get(no, []) if _canon(x)}
        c_canon = {_canon(x): x for x in cur_names.get(no, []) if _canon(x)}
        if not s_canon and not c_canon:
            continue
        common = [k for k in s_canon if k in c_canon]
        snap_only = [k for k in s_canon if k not in c_canon]
        cur_only = [k for k in c_canon if k not in s_canon]
        disp = {**c_canon, **s_canon}  # 표기는 스냅샷 우선

        # 기대사업명 베이스 = 공통 + 스냅샷only (현재 독립). cur_only 는 검토 칸으로.
        base = [disp[k] for k in common] + [disp[k] for k in snap_only]
        blob = cur_blob.get(no, "")
        review_lines: List[str] = []
        for k in cur_only:
            nm = c_canon[k]
            grounded = _name_grounded(nm, blob) if blob else True
            review_lines.append(f"{nm}  [{'근거O' if grounded else '근거X·환각의심'}]")
            if not grounded:
                n_ungrounded += 1
        n_review += len(cur_only)

        summary = f"공통 {len(common)} · 스냅샷only {len(snap_only)} · 현재only {len(cur_only)}"
        ws.append([no, q, "\n".join(base), "\n".join(review_lines), summary, "", ""])
        rr = ws.max_row
        ws.cell(row=rr, column=3).fill = EDIT_FILL
        if review_lines:
            ws.cell(row=rr, column=4).fill = REVIEW_FILL
        for cc in (2, 3, 4, 7):
            ws.cell(row=rr, column=cc).alignment = WRAP
        n_rows += 1

    last = ws.max_row
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLS))}{last}"
    dv = DataValidation(type="list", formula1='"미검토,검토완료,제외"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{last}")

    g = wb.create_sheet("사용법", 0)
    g.column_dimensions["A"].width = 106
    for line in [
        "■ 사업명 골든셋 후보 — 사람 1패스 확정용",
        "",
        "· '기대사업명' = 채점 기준. 프리필 = 공통(두출처 일치) + 스냅샷only → 현재 출력과 독립적인 베이스.",
        "· '현재추가후보' = 현재 답변에만 있던 사업명(= 정밀도를 좌우). 한 줄씩 [근거O/X] 표시:",
        "    - 채택할 것 → '기대사업명' 칸으로 옮겨 적기(이 질문에 정말 맞는 사업).",
        "    - 버릴 것 → 그냥 둠. 특히 [근거X·환각의심](검색문서에 근거 없음)은 대체로 버림.",
        "· 즉 사람 작업은 '현재추가후보 검토'에 집중하면 됨(공통/스냅샷 베이스는 대체로 유지).",
        "· 비추천형(단일조회·비교)으로 기대 사업명이 없으면 '기대사업명' 비우고 검토=제외.",
        "· 연도·지역 접두어는 채점 시 자동 정규화.",
        "",
        "⚠ 현재 출력에서 파생한 gold로는 정확도가 편향됨(union→정밀도 만점). 그래서 현재only를",
        "   베이스에서 빼고 '검토 후 채택'으로 둔 것 — 사람이 확정해야 진짜 P/R이 나온다.",
        "",
        f"■ 채점: python tests/score_against_gold.py {out.name} <consistency.jsonl>",
    ]:
        g.append([line])
    g["A1"].font = Font(bold=True, size=13)

    wb.save(out)
    print(f"[done] gold 후보(검토최적화) → {out}")
    print(f"  문항 {n_rows} | 현재only 검토항목 {n_review}개 (그중 환각의심[근거X] {n_ungrounded}개)")
    print("  '현재추가후보' 칸만 검토해 채택분을 '기대사업명'으로 올린 뒤 채점하세요.")


if __name__ == "__main__":
    main()
