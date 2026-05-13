"""260512_gsnd_sample_set.xlsx 의 ver1/ver2 사용자 입력에 대한 답변셋 생성.

- 두 시트의 '사용자 입력'은 동일하므로 질문당 한 번씩만 챗 API 호출.
- 봇이 시군 되묻기를 하면 행의 '지역' 컬럼 값을 후속 입력으로 전달
  ('A→B' 이사 시나리오는 'B', 'A/B' 다중 시군은 'A' 우선, 비어있으면 '창원' 폴백).
- 결과는 봇 되묻기 / 추가 입력 / 봇 최종답변 / 실제 검색된 문서 컬럼에 채워
  tests/llm_eval_results/gsnd_sample_set_<timestamp>.xlsx 로 저장.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
import time
import uuid
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parent
XLSX_IN = ROOT / "260512_gsnd_sample_set.xlsx"
OUT_DIR = ROOT / "llm_eval_results"
BASE_URL = "http://localhost:8000"
ENDPOINT = f"{BASE_URL}/v1/chat/completions"
TIMEOUT = 300.0
FALLBACK_SIGUN = "창원"
RETRY_MAX = 3
RETRY_BACKOFF_SEC = 5.0

# 시트 컬럼 인덱스 (0-based)
COL_TYPE = 0
COL_DIFF = 1
COL_REGION = 2
COL_USER_INPUT = 3
COL_BOT_REASK = 4
COL_FOLLOWUP_INPUT = 5
COL_BOT_FINAL = 6
COL_EXPECTED_DOC = 7
COL_RETRIEVED_DOC = 8


def pick_followup_region(region_cell: str | None) -> str:
    """지역 컬럼 → 후속 입력으로 사용할 시군명."""
    region = (region_cell or "").strip()
    if not region:
        return FALLBACK_SIGUN
    if "→" in region:
        return region.split("→", 1)[1].strip() or FALLBACK_SIGUN
    if "/" in region:
        return region.split("/", 1)[0].strip() or FALLBACK_SIGUN
    return region


def is_sigun_clarify(text: str) -> bool:
    t = text or ""
    return ("시/군" in t or "시군" in t) and ("거주" in t or "지역" in t or "알려" in t)


def post_chat(client: httpx.Client, messages: list[dict], user_id: str, conv_id: str) -> dict:
    payload = {
        "messages": messages,
        "user_id": user_id,
        "conv_id": conv_id,
        "stream": False,
    }
    last_err: Exception | None = None
    for attempt in range(1, RETRY_MAX + 1):
        try:
            r = client.post(ENDPOINT, json=payload, timeout=TIMEOUT)
            r.raise_for_status()
            return r.json()
        except (httpx.RemoteProtocolError, httpx.ReadError, httpx.ConnectError,
                httpx.ReadTimeout, httpx.WriteError) as e:
            last_err = e
            if attempt < RETRY_MAX:
                wait = RETRY_BACKOFF_SEC * attempt
                print(f"    [retry {attempt}/{RETRY_MAX - 1}] {type(e).__name__}: {e} → {wait}s 후 재시도")
                time.sleep(wait)
                continue
            raise
    raise last_err  # type: ignore[misc]


def extract_content(resp: dict) -> tuple[str, bool]:
    try:
        content = resp["choices"][0]["message"]["content"]
    except Exception:
        content = json.dumps(resp, ensure_ascii=False)[:500]
    return content, bool(resp.get("is_clarification"))


_DOC_NAME_PATTERN = re.compile(r"^(?:20\d{2}_)?(?P<sigun>[^_]+)_(?P<title>.+?)(?:\.[A-Za-z0-9]+)?$")


def _format_doc_name(raw: str) -> str:
    """'2026_양산시_아동수당.hwpx' → '아동수당 (양산시)'"""
    raw = raw.strip()
    if not raw:
        return raw
    m = _DOC_NAME_PATTERN.match(raw)
    if m:
        return f"{m.group('title')} ({m.group('sigun')})"
    if "." in raw:
        raw = raw.rsplit(".", 1)[0]
    return raw


def extract_referenced(resp: dict) -> str:
    refs = resp.get("referenced_documents") or []
    if not refs:
        return ""
    lines: list[str] = []
    seen: set[str] = set()
    for d in refs:
        name = str(d.get("name") or "").strip()
        doc_id = str(d.get("id") or "").strip()
        label = _format_doc_name(name) if name else (f"id={doc_id}" if doc_id else "")
        if label and label not in seen:
            seen.add(label)
            lines.append(label)
    return "\n".join(lines)


def run_one(client: httpx.Client, question: str, region_cell: str) -> dict:
    user_id = f"qa_set_{uuid.uuid4().hex[:8]}"
    conv_id = str(uuid.uuid4())
    msgs: list[dict] = [{"role": "user", "content": question}]

    t0 = time.monotonic()
    try:
        resp1 = post_chat(client, msgs, user_id, conv_id)
    except Exception as e:
        return {
            "question": question,
            "bot_reask": "",
            "followup_input": "",
            "final_answer": f"[ERROR] {e}",
            "referenced_docs": "",
            "elapsed_sec": round(time.monotonic() - t0, 2),
        }

    ans1, is_clar = extract_content(resp1)
    refs1 = extract_referenced(resp1)
    bot_reask = ""
    followup_input = ""
    final_answer = ans1
    referenced = refs1

    if is_clar and is_sigun_clarify(ans1):
        bot_reask = ans1
        followup_input = pick_followup_region(region_cell)
        msgs.append({"role": "assistant", "content": ans1})
        msgs.append({"role": "user", "content": followup_input})
        try:
            resp2 = post_chat(client, msgs, user_id, conv_id)
            final_answer, _ = extract_content(resp2)
            referenced = extract_referenced(resp2) or refs1
        except Exception as e:
            final_answer = f"[ERROR follow-up] {e}"

    return {
        "question": question,
        "bot_reask": bot_reask,
        "followup_input": followup_input,
        "final_answer": final_answer,
        "referenced_docs": referenced,
        "elapsed_sec": round(time.monotonic() - t0, 2),
    }


def collect_rows(wb: openpyxl.Workbook) -> dict[str, list[tuple[int, str, str]]]:
    """sheet_name -> [(row_idx, question, region)] (헤더 제외)."""
    out: dict[str, list[tuple[int, str, str]]] = {}
    for sheet_name in ("ver1", "ver2"):
        ws = wb[sheet_name]
        rows: list[tuple[int, str, str]] = []
        for row_idx in range(2, ws.max_row + 1):
            q = ws.cell(row=row_idx, column=COL_USER_INPUT + 1).value
            r = ws.cell(row=row_idx, column=COL_REGION + 1).value
            if not q:
                continue
            rows.append((row_idx, str(q).strip(), str(r or "").strip()))
        out[sheet_name] = rows
    return out


def main(limit: int | None = None) -> None:
    if not XLSX_IN.exists():
        print(f"[error] 입력 파일이 없습니다: {XLSX_IN}")
        sys.exit(1)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_out = OUT_DIR / f"gsnd_sample_set_{ts}.xlsx"
    jsonl_out = OUT_DIR / f"gsnd_sample_set_{ts}.jsonl"

    wb = openpyxl.load_workbook(XLSX_IN)
    sheet_rows = collect_rows(wb)

    # ver1·ver2 질문이 동일하므로 캐시로 1회만 호출.
    question_cache: dict[str, dict] = {}

    # 처리할 (질문, 대표 지역) 목록 — 첫 등장 행의 지역값 사용.
    work_items: list[tuple[str, str]] = []
    seen: set[str] = set()
    for sheet_name in ("ver1", "ver2"):
        for _, q, r in sheet_rows[sheet_name]:
            if q in seen:
                continue
            seen.add(q)
            work_items.append((q, r))
    if limit:
        work_items = work_items[:limit]

    print(f"[run] 고유 질문 {len(work_items)}개 처리 시작 → {xlsx_out}")

    with httpx.Client() as client, jsonl_out.open("w", encoding="utf-8") as fjl:
        for i, (q, region) in enumerate(work_items, 1):
            result = run_one(client, q, region)
            question_cache[q] = result
            fjl.write(json.dumps(result, ensure_ascii=False) + "\n")
            fjl.flush()
            preview = (result["final_answer"] or "")[:100].replace("\n", " ")
            tag = " (clarified)" if result["bot_reask"] else ""
            print(f"[{i}/{len(work_items)}] {result['elapsed_sec']}s{tag} | {preview}")

    # 두 시트 모두에 결과 채우기.
    for sheet_name in ("ver1", "ver2"):
        ws = wb[sheet_name]
        for row_idx, q, _ in sheet_rows[sheet_name]:
            r = question_cache.get(q)
            if not r:
                continue
            ws.cell(row=row_idx, column=COL_BOT_REASK + 1, value=r["bot_reask"])
            ws.cell(row=row_idx, column=COL_FOLLOWUP_INPUT + 1, value=r["followup_input"])
            ws.cell(row=row_idx, column=COL_BOT_FINAL + 1, value=r["final_answer"])
            ws.cell(row=row_idx, column=COL_RETRIEVED_DOC + 1, value=r["referenced_docs"])

    wb.save(xlsx_out)
    print(f"[done] xlsx={xlsx_out}\n[done] jsonl={jsonl_out}")


if __name__ == "__main__":
    n = int(sys.argv[1]) if len(sys.argv) > 1 else None
    main(n)
