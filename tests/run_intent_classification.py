"""qa_test/260513_gsnd_total_dataset_v0.3.xlsx 의 question 컬럼에 대해
unified_preprocess() 를 직접 호출하여 intent 분류 결과를 dump.

- HTTP 서버 불필요. 단 sLLM(8B) / 32B LLM 서버는 떠 있어야 함
  (call_classifier_with_fallback 가 호출).
- use_rag=False 로 호출 → keyword 추출/쿼리 확장 스킵, 순수 intent 분류만 빠르게.
- 결과는 qa_test/intent_classification_<timestamp>.xlsx + .jsonl 두 형식.

사용:
    cd backend-qa
    python tests/run_intent_classification.py           # 전체
    python tests/run_intent_classification.py 5         # 처음 5개 (smoke test)
"""
from __future__ import annotations

import asyncio
import datetime as dt
import json
import sys
import time
from pathlib import Path


def _ensure_app_path_on_sys() -> None:
    root = Path(__file__).resolve().parents[1]
    src = root / "src"
    if str(src) not in sys.path:
        sys.path.insert(0, str(src))


_ensure_app_path_on_sys()

import openpyxl  # noqa: E402

import app.chat.infra.rag  # noqa: E402, F401  — 순환 import 회피 pre-import
from app.chat.preprocessing import unified_preprocess  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
XLSX_IN = ROOT / "qa_test" / "260513_gsnd_total_dataset_v0.3.xlsx"
OUT_DIR = ROOT / "qa_test"

HEADERS = [
    "no", "question",
    "intent", "intent_reason",
    "search_target", "policy_priority_tag", "detail_requested",
    "reformed_query",
    "elapsed_sec", "error",
]


def load_questions() -> list[tuple[int, str]]:
    wb = openpyxl.load_workbook(XLSX_IN, read_only=True, data_only=True)
    ws = wb["gsnd_total"] if "gsnd_total" in wb.sheetnames else wb.active
    out: list[tuple[int, str]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None or r[1] is None:
            continue
        try:
            no = int(r[0])
        except (TypeError, ValueError):
            continue
        q = str(r[1]).strip()
        if q:
            out.append((no, q))
    return out


async def run_one(no: int, question: str) -> dict:
    t0 = time.monotonic()
    try:
        r = await unified_preprocess(user_query=question, messages=[], use_rag=False)
        err = ""
    except Exception as e:
        r = {}
        err = f"{type(e).__name__}: {e}"
    elapsed = round(time.monotonic() - t0, 2)
    return {
        "no": no, "question": question,
        "intent": r.get("intent", ""),
        "intent_reason": r.get("intent_reason", ""),
        "search_target": r.get("search_target") or "",
        "policy_priority_tag": r.get("policy_priority_tag") or "",
        "detail_requested": r.get("detail_requested", ""),
        "reformed_query": r.get("reformed_query", ""),
        "elapsed_sec": elapsed,
        "error": err,
    }


async def main(limit: int | None = None) -> None:
    if not XLSX_IN.exists():
        print(f"[error] 입력 파일 없음: {XLSX_IN}")
        sys.exit(1)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    questions = load_questions()
    if limit:
        questions = questions[:limit]

    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = OUT_DIR / f"intent_classification_{ts}.xlsx"
    out_jsonl = OUT_DIR / f"intent_classification_{ts}.jsonl"

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "intent"
    ws_out.append(HEADERS)

    print(f"[run] {len(questions)} questions → xlsx={out_xlsx}")

    with out_jsonl.open("w", encoding="utf-8") as fj:
        for i, (no, q) in enumerate(questions, 1):
            r = await run_one(no, q)
            ws_out.append([r[h] for h in HEADERS])
            wb_out.save(out_xlsx)  # 매 행마다 저장 (장시간 실행 대비)
            fj.write(json.dumps(r, ensure_ascii=False) + "\n")
            fj.flush()
            tag = f" ERR={r['error'][:60]}" if r["error"] else ""
            preview = q[:50].replace("\n", " ")
            print(
                f"[{i}/{len(questions)}] no={no} intent={r['intent'] or '-':<16}"
                f" {r['elapsed_sec']}s | {preview}{tag}"
            )

    print(f"[done] xlsx={out_xlsx}")
    print(f"[done] jsonl={out_jsonl}")


if __name__ == "__main__":
    n = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1] not in ("", "0") else None
    asyncio.run(main(n))
