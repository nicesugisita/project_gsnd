# -*- coding: utf-8 -*-
"""싱글턴 테스트셋 러너 — 광역 질문은 칩(버튼) 드릴다운까지 자동 재현해 카드를 캡처.

흐름(질문 1건당):
  [1] 질문 전송(plain /v1/chat/completions, mode 미지정 → 실제 라우팅 그대로).
      └ 시군 되묻기(is_clarification + "시/군")면 follow_up(기본 '창원')으로 후속 1회.
  [2] 추천 응답에서 topic_chips / slots 수신.
      - 칩 있음(광역=direct_topic): 각 칩을 slots+label 로 drilldown 재전송 →
        guide_services(서비스 카드) 캡처. ★결정적·LLM 0회(서버가 pre_check/재작성/분류 생략).
      - 칩 없음(소량=llm): 자연어 답변 + referenced_documents 가 곧 결과.

전제: 서버가 떠 있어야 함(기본 127.0.0.1:8000). guide_recommend DB 도달 필요.

사용:
  python tests/run_singleturn_drilldown.py                      # sheet1 150건 전체, 모든 칩 드릴다운
  python tests/run_singleturn_drilldown.py --limit 2            # 앞 2건 스모크
  python tests/run_singleturn_drilldown.py --no 7               # 7번만
  python tests/run_singleturn_drilldown.py --start-no 50        # 50번부터 resume
  python tests/run_singleturn_drilldown.py --max-chips 3        # 칩 상위 3개만(기본 0=전체)
  python tests/run_singleturn_drilldown.py --src <xlsx>         # 다른 입력(시트 gsnd_total / no·question·clarified·follow_up)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import time
import uuid
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = ROOT / "tests" / "integrated_split" / "sheet1_dqp_single.xlsx"
OUT_DIR = ROOT / "qa_test"
TIMEOUT = 300.0
RETRY_MAX = 3
RETRY_BACKOFF = 5.0
USER_PREFIX = "qa_st_dd"
DEFAULT_SIGUN = "창원"


def _iter_rows(src: Path):
    """시트 gsnd_total: no / question / clarified / follow_up."""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb["gsnd_total"] if "gsnd_total" in wb.sheetnames else wb.worksheets[0]
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None or r[1] is None:
            continue
        try:
            no = int(r[0])
        except (TypeError, ValueError):
            continue
        q = str(r[1]).strip()
        if not q:
            continue
        clarified = (str(r[2]).strip().upper() == "Y") if len(r) > 2 and r[2] else False
        follow_up = str(r[3]).strip() if len(r) > 3 and r[3] else ""
        yield {"no": no, "question": q, "clarified": clarified, "follow_up": follow_up}
    wb.close()


def _is_sigun_clarify(text: str) -> bool:
    t = text or ""
    return ("시/군" in t or "시군" in t) and ("거주" in t or "지역" in t or "알려" in t)


def _post(client: httpx.Client, endpoint: str, *, messages, conv_id, user_id, drilldown=None) -> dict:
    payload = {"messages": messages, "user_id": user_id, "conv_id": conv_id, "stream": False}
    if drilldown is not None:
        payload["drilldown"] = drilldown
    last = None
    for attempt in range(1, RETRY_MAX + 1):
        try:
            r = client.post(endpoint, json=payload, timeout=TIMEOUT)
            if 500 <= r.status_code < 600 and attempt < RETRY_MAX:
                time.sleep(RETRY_BACKOFF * attempt)
                continue
            r.raise_for_status()
            return r.json()
        except (httpx.RemoteProtocolError, httpx.ReadError, httpx.ConnectError,
                httpx.ReadTimeout, httpx.WriteError) as e:
            last = e
            if attempt < RETRY_MAX:
                time.sleep(RETRY_BACKOFF * attempt)
                continue
            raise
    raise last  # type: ignore[misc]


def _content(resp: dict) -> str:
    return resp.get("choices", [{}])[0].get("message", {}).get("content", "") or ""


def _run_one(client, endpoint, row: dict, max_chips: int) -> dict:
    """질문 1건 → 추천 + (광역이면) 칩 전체 드릴다운 결과."""
    no, question = row["no"], row["question"]
    conv_id = f"{USER_PREFIX}-{no}-{uuid.uuid4().hex[:6]}"
    msgs = [{"role": "user", "content": question}]

    resp = _post(client, endpoint, messages=msgs, conv_id=conv_id, user_id=conv_id)
    sigun_used = ""
    # 시군 되묻기 → follow_up(기본 창원)으로 후속 1회
    if bool(resp.get("is_clarification")) and _is_sigun_clarify(_content(resp)):
        nxt = row["follow_up"] or (DEFAULT_SIGUN if row["clarified"] else DEFAULT_SIGUN)
        sigun_used = nxt
        msgs.append({"role": "assistant", "content": _content(resp)})
        msgs.append({"role": "user", "content": nxt})
        resp = _post(client, endpoint, messages=msgs, conv_id=conv_id, user_id=conv_id)

    header = _content(resp)
    chips = resp.get("topic_chips") or []
    slots = resp.get("slots") or {}
    refs = resp.get("referenced_documents") or []
    first_cards = resp.get("guide_services") or []
    mode = "broad" if chips else ("cards" if first_cards else "narrow")
    if not sigun_used:
        sigun_used = (slots.get("sigun") or "").replace("경상남도", "").strip()

    drills = []
    if chips:
        sel = chips if max_chips <= 0 else chips[:max_chips]
        for c in sel:
            label = c.get("label")
            if not label:
                continue
            dd = {
                "sigun": slots.get("sigun"),
                "lifecycle": slots.get("lifecycle", []),
                "household": slots.get("household", []),
                "topic_category": [label],
            }
            dresp = _post(client, endpoint, messages=[{"role": "user", "content": label}],
                          conv_id=conv_id, user_id=conv_id, drilldown=dd)
            cards = dresp.get("guide_services") or []
            drills.append({
                "label": label,
                "chip_count": c.get("count"),
                "header": _content(dresp),
                "n_cards": len(cards),
                "cards": [{k: g.get(k) for k in ("service_name", "target", "benefit")} for g in cards],
            })

    return {
        "no": no, "question": question, "conv_id": conv_id, "sigun": sigun_used,
        "mode": mode, "header": header,
        "chips": [{"label": c.get("label"), "count": c.get("count")} for c in chips],
        "n_chips": len(chips),
        "narrow_refs": [d.get("name") for d in refs] if mode == "narrow" else [],
        "first_cards": [g.get("service_name") for g in first_cards],
        "drilldowns": drills,
        "n_cards_total": sum(d["n_cards"] for d in drills) + len(first_cards),
    }


def _write_xlsx(records: list[dict], out: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["no", "question", "sigun", "mode", "n_chips",
               "chips(label:count)", "n_cards_total", "header"])
    for r in records:
        chips_s = "; ".join(f"{c['label']}:{c['count']}" for c in r["chips"])
        ws.append([r["no"], r["question"], r["sigun"], r["mode"], r["n_chips"],
                   chips_s, r["n_cards_total"], (r["header"] or "")[:200]])

    wc = wb.create_sheet("cards")
    wc.append(["no", "question", "source", "chip_count", "rank",
               "service_name", "target", "benefit"])
    for r in records:
        if r["mode"] == "narrow":
            for i, nm in enumerate(r["narrow_refs"], 1):
                wc.append([r["no"], r["question"], "narrow", "", i, nm, "", ""])
        for d in r["drilldowns"]:
            for i, c in enumerate(d["cards"], 1):
                wc.append([r["no"], r["question"], d["label"], d["chip_count"], i,
                           c.get("service_name"), c.get("target"),
                           (c.get("benefit") or "")[:120]])
    wb.save(out)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(DEFAULT_SRC))
    ap.add_argument("--endpoint", default="http://127.0.0.1:8000/v1/chat/completions")
    ap.add_argument("--no", type=int, default=None, help="특정 no만")
    ap.add_argument("--start-no", type=int, default=1, help="이 no 이상부터(resume)")
    ap.add_argument("--limit", type=int, default=None, help="앞 N건만")
    ap.add_argument("--max-chips", type=int, default=0, help="칩 드릴다운 상한(0=전체)")
    ap.add_argument("--tag", default=None)
    args = ap.parse_args()

    src = Path(args.src)
    if not src.exists():
        print(f"[error] 입력 없음: {src}"); sys.exit(1)

    rows = list(_iter_rows(src))
    if args.no is not None:
        rows = [r for r in rows if r["no"] == args.no]
    else:
        rows = [r for r in rows if r["no"] >= args.start_no]
        if args.limit:
            rows = rows[: args.limit]
    if not rows:
        print("[error] 대상 질문 0개"); sys.exit(1)

    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    tag = f"{args.tag}_" if args.tag else ""
    cap_jsonl = OUT_DIR / f"singleturn_dd_{tag}{ts}.jsonl"
    out_xlsx = OUT_DIR / f"singleturn_dd_{tag}{ts}.xlsx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"[run] 질문 {len(rows)}건 | max_chips={'전체' if args.max_chips<=0 else args.max_chips} | endpoint={args.endpoint}")
    records = []
    with httpx.Client() as client, cap_jsonl.open("w", encoding="utf-8") as fj:
        for i, row in enumerate(rows, 1):
            t0 = time.monotonic()
            try:
                rec = _run_one(client, args.endpoint, row, args.max_chips)
                err = "-"
            except Exception as e:
                rec = {"no": row["no"], "question": row["question"], "sigun": "", "mode": "ERROR",
                       "header": f"{type(e).__name__}: {e}", "chips": [], "n_chips": 0,
                       "narrow_refs": [], "first_cards": [], "drilldowns": [], "n_cards_total": 0}
                err = f"{type(e).__name__}: {e}"
            records.append(rec)
            fj.write(json.dumps(rec, ensure_ascii=False) + "\n"); fj.flush()
            el = round(time.monotonic() - t0, 1)
            print(f"[{i}/{len(rows)}] no={row['no']} {el}s mode={rec['mode']} "
                  f"chips={rec['n_chips']} cards={rec['n_cards_total']} err={err} :: {row['question'][:24]}")

    _write_xlsx(records, out_xlsx)
    print(f"\n[done] jsonl → {cap_jsonl}")
    print(f"[done] xlsx  → {out_xlsx} (summary / cards 시트)")


if __name__ == "__main__":
    main()
