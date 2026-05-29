"""답변 정확도 LLM-as-judge 채점 — answer_v2(최종응답) vs 정답(final_answer) → 문항별 정·부분·오답.

정답 기준: 260513_gsnd_total_dataset_v0.3.xlsx 의 gsnd_total 시트 `final_answer` 컬럼(사람 확정 모범답안).
평가 대상: 260527_answer_v2.xlsx 의 stages 시트 `최종응답(response)` (질문당 회차 여럿 → R1 = 최저 round).
판정자: 로컬 32B (.env 의 LLM_API_URL/MODEL_NAME). temperature=0, seed 고정 → 가능한 한 결정적.

루브릭(JSON 반환):
- factual_accuracy 0~2 : 금액·대상·조건·제도명 등 사실이 정답과 부합(허위·과장 감점).
- coverage         0~2 : 정답이 담은 핵심 제도/정보를 빠짐없이 포함.
- relevance        0~1 : 질문 의도에 맞게 답함(엉뚱한 제도로 새지 않음).
- 형식(구조화 vs 산문)·말투 차이는 감점하지 않는다. **내용 기준**.
- verdict : 정답 / 부분정답 / 오답.
정확도 = verdict 점수(정답1.0·부분정답0.5·오답0)의 평균 + 차원별 Macro 평균.

사용:
    python tests/judge_answer_accuracy.py                 # 기본 경로, 전체(매칭 문항)
    python tests/judge_answer_accuracy.py --limit 5       # 스모크(앞 5문항만)
    python tests/judge_answer_accuracy.py --answer <a.xlsx> --gold <g.xlsx> --out <o.xlsx>
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tests"))
from build_review_checklist import _norm  # noqa: E402

DEF_ANSWER = ROOT / "qa_test" / "260527_answer_v2.xlsx"
DEF_GOLD = ROOT / "qa_test" / "260513_gsnd_total_dataset_v0.3.xlsx"

# stages 컬럼 인덱스(0-base): 질문0, round2, 최종응답17
C_Q, C_ROUND, C_RESP = 0, 2, 17

VERDICT_SCORE = {"정답": 1.0, "부분정답": 0.5, "오답": 0.0}
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="E7E6E6")
HDR = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
_CELL_MAX = 4000

SYSTEM_PROMPT = (
    "당신은 한국 지자체 복지 상담 챗봇의 답변을 채점하는 엄정한 평가자다. "
    "주어진 '정답(모범답안)'을 기준으로 '평가대상 답변'의 내용 정확성을 판정한다. "
    "형식(서비스 나열 구조 vs 산문체)이나 말투 차이는 절대 감점하지 않고, 오직 내용으로만 판단한다. "
    "반드시 JSON 객체 하나만 출력하고 그 외 설명·머리말은 출력하지 않는다."
)

USER_TMPL = """[질문]
{question}

[정답(모범답안)]
{gold}

[평가대상 답변]
{cand}

[채점 기준]
- factual_accuracy (0~2): 금액·지원대상·조건·제도명 등 사실이 정답과 부합하는가. 정답에 없는 허위·과장 정보는 감점.
- coverage (0~2): 정답이 담은 핵심 제도/정보를 빠짐없이 포함했는가(누락 시 감점).
- relevance (0~1): 질문 의도에 맞게 답했는가(무관한 제도로 새면 0).
- verdict: 위 종합 판정. "정답"(핵심 사실 정확·누락 거의 없음) / "부분정답"(일부 맞으나 누락 또는 일부 오류) / "오답"(핵심이 틀리거나 질문과 무관).

아래 형식의 JSON 하나만 출력하라:
{{"factual_accuracy": <0-2 정수>, "coverage": <0-2 정수>, "relevance": <0-1 정수>, "verdict": "정답|부분정답|오답", "reason": "한두 문장 근거"}}"""


def _trunc(s: Any) -> str:
    s = str(s or "")
    return s if len(s) <= _CELL_MAX else s[:_CELL_MAX] + f"...[+{len(s) - _CELL_MAX}]"


def _load_env() -> Dict[str, str]:
    """레포 루트의 .env* 파일에서 KEY=VALUE 로드 (os.environ 우선)."""
    import os
    env: Dict[str, str] = {}
    for p in sorted(ROOT.glob(".env*")):
        if not p.is_file():
            continue
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    env.update({k: v for k, v in os.environ.items() if k in ("LLM_API_URL", "MODEL_NAME", "FIXED_LLM_SEED")})
    return env


def _load_gold(path: Path) -> Dict[str, str]:
    """norm(question) → final_answer (채워진 행만)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["gsnd_total"]
    out: Dict[str, str] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[1] is None or len(r) <= 4 or not r[4]:
            continue
        out.setdefault(_norm(r[1]), str(r[4]).strip())
    wb.close()
    return out


def _load_answers(path: Path) -> Dict[str, Dict[str, Any]]:
    """norm(question) → {질문, 응답(R1=최저 round)}."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["stages"]
    best: Dict[str, Dict[str, Any]] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[C_Q] is None:
            continue
        k = _norm(r[C_Q])
        rnd = r[C_ROUND] if isinstance(r[C_ROUND], int) else 0
        if k not in best or rnd < best[k]["round"]:
            best[k] = {"question": str(r[C_Q]), "round": rnd, "resp": str(r[C_RESP] or "")}
    wb.close()
    return best


_JSON_RE = re.compile(r"\{.*\}", re.DOTALL)


def _parse_judge(text: str) -> Optional[Dict[str, Any]]:
    """LLM 응답에서 JSON 객체 추출·파싱."""
    if not text:
        return None
    m = _JSON_RE.search(text)
    if not m:
        return None
    try:
        d = json.loads(m.group(0))
    except json.JSONDecodeError:
        return None
    v = str(d.get("verdict", "")).strip()
    if v not in VERDICT_SCORE:
        return None

    def _int(x, lo, hi):
        try:
            return max(lo, min(hi, int(round(float(x)))))
        except (TypeError, ValueError):
            return 0
    return {
        "factual_accuracy": _int(d.get("factual_accuracy"), 0, 2),
        "coverage": _int(d.get("coverage"), 0, 2),
        "relevance": _int(d.get("relevance"), 0, 1),
        "verdict": v,
        "reason": str(d.get("reason", "")).strip(),
    }


def _judge(client: httpx.Client, url: str, model: str, seed: int,
           question: str, gold: str, cand: str) -> Dict[str, Any]:
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": USER_TMPL.format(question=question, gold=gold, cand=cand)},
        ],
        "temperature": 0,
        "top_p": 1,
        "seed": seed,
        "max_tokens": 512,
        "stream": False,
    }
    last_err = ""
    for attempt in range(3):
        try:
            # 1차는 JSON 강제, 거부(400)하면 자유형식으로 폴백
            body = dict(payload)
            if attempt == 0:
                body["response_format"] = {"type": "json_object"}
            resp = client.post(url, json=body, timeout=120)
            if resp.status_code == 400 and attempt == 0:
                continue  # response_format 미지원 → 폴백
            resp.raise_for_status()
            content = resp.json().get("choices", [{}])[0].get("message", {}).get("content", "")
            parsed = _parse_judge(content)
            if parsed:
                return parsed
            last_err = f"파싱실패: {content[:120]}"
        except Exception as e:  # noqa: BLE001
            last_err = f"{type(e).__name__}: {e}"
            time.sleep(2)
    return {"factual_accuracy": "", "coverage": "", "relevance": "", "verdict": "ERROR", "reason": last_err}


_COLS = [
    ("no", 5), ("질문", 38), ("판정", 9), ("정확도점수", 9),
    ("사실정확(0-2)", 11), ("커버리지(0-2)", 11), ("관련성(0-1)", 10), ("정규화점수", 10),
    ("근거", 50), ("정답(모범답안)", 50), ("모델 최종응답(R1)", 50),
]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--answer", default=str(DEF_ANSWER))
    ap.add_argument("--gold", default=str(DEF_GOLD))
    ap.add_argument("--out", default=None)
    ap.add_argument("--limit", type=int, default=None, help="앞 N문항만(스모크)")
    args = ap.parse_args()

    ans_p, gold_p = Path(args.answer), Path(args.gold)
    out = Path(args.out) if args.out else ans_p.with_name(ans_p.stem + "_accuracy.xlsx")
    env = _load_env()
    url = (env.get("LLM_API_URL") or "").rstrip("/") + "/v1/chat/completions"
    model = env.get("MODEL_NAME") or ""
    seed = int(env.get("FIXED_LLM_SEED") or 42)
    if not env.get("LLM_API_URL") or not model:
        print("[error] .env 에서 LLM_API_URL / MODEL_NAME 을 찾지 못함"); sys.exit(1)

    gold = _load_gold(gold_p)
    ans = _load_answers(ans_p)
    keys = sorted(set(gold) & set(ans), key=lambda k: ans[k]["question"])
    if args.limit:
        keys = keys[: args.limit]
    print(f"[info] 정답 {len(gold)}종 | 응답 {len(ans)}종 | 채점대상 {len(keys)}문항")
    print(f"[info] judge LLM = {model} @ {url} (seed={seed})")
    if not keys:
        print("[error] 채점할 문항 0개 (질문 매칭 확인)"); sys.exit(1)

    rows: List[Dict[str, Any]] = []
    with httpx.Client() as client:
        for i, k in enumerate(keys, 1):
            q = ans[k]["question"]
            cand = ans[k]["resp"]
            g = gold[k]
            t0 = time.monotonic()
            j = _judge(client, url, model, seed, q, g, cand)
            el = round(time.monotonic() - t0, 1)
            vs = VERDICT_SCORE.get(j["verdict"], "")
            norm_score = ""
            if isinstance(j["factual_accuracy"], int):
                norm_score = round((j["factual_accuracy"] + j["coverage"] + j["relevance"]) / 5, 2)
            rows.append({
                "no": i, "질문": q, "판정": j["verdict"], "정확도점수": vs,
                "사실정확(0-2)": j["factual_accuracy"], "커버리지(0-2)": j["coverage"],
                "관련성(0-1)": j["relevance"], "정규화점수": norm_score,
                "근거": j["reason"], "정답(모범답안)": _trunc(g), "모델 최종응답(R1)": _trunc(cand),
            })
            print(f"[{i}/{len(keys)}] {el}s {j['verdict']:>4} (score={vs}) :: {q[:30]}")

    # ── xlsx 작성 ─────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "채점"
    ws.append([c[0] for c in _COLS])
    for j_, c in enumerate(_COLS, 1):
        cell = ws.cell(row=1, column=j_)
        cell.font = HDR_FONT; cell.fill = HDR; cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(j_)].width = c[1]
    for row in rows:
        ws.append([row[c[0]] for c in _COLS])
        rr = ws.max_row
        vc = ws.cell(row=rr, column=3)
        vc.fill = {"정답": GREEN, "부분정답": YELLOW, "오답": RED}.get(row["판정"], GREY)
        for ci in (2, 9, 10, 11):
            ws.cell(row=rr, column=ci).alignment = WRAP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLS))}{ws.max_row}"

    # 집계
    scored = [r for r in rows if isinstance(r["정확도점수"], float)]
    n = len(scored)
    dist = defaultdict(int)
    for r in rows:
        dist[r["판정"]] += 1
    acc = sum(r["정확도점수"] for r in scored) / n if n else 0.0
    mf = sum(r["사실정확(0-2)"] for r in scored) / n if n else 0.0
    mc = sum(r["커버리지(0-2)"] for r in scored) / n if n else 0.0
    mr = sum(r["관련성(0-1)"] for r in scored) / n if n else 0.0

    s = wb.create_sheet("요약", 0)
    s.column_dimensions["A"].width = 28; s.column_dimensions["B"].width = 14
    s.append(["답변 정확도 LLM 채점 요약"]); s["A1"].font = Font(bold=True, size=13)
    s.append(["judge LLM", model])
    s.append(["채점 문항", len(rows)])
    s.append(["유효 채점(ERROR 제외)", n]); s.append([])
    s.append(["정확도(정답1·부분0.5·오답0 평균)", round(acc, 3)]); s.cell(row=s.max_row, column=1).font = HDR_FONT
    s.append([]); s.append(["[판정 분포]"]); s.cell(row=s.max_row, column=1).font = HDR_FONT
    for v in ("정답", "부분정답", "오답", "ERROR"):
        if dist.get(v):
            s.append([v, dist[v], round(dist[v] / len(rows), 3)])
            s.cell(row=s.max_row, column=3).number_format = "0%"
    s.append([]); s.append(["[차원별 Macro 평균]"]); s.cell(row=s.max_row, column=1).font = HDR_FONT
    s.append(["사실정확(0-2)", round(mf, 2)])
    s.append(["커버리지(0-2)", round(mc, 2)])
    s.append(["관련성(0-1)", round(mr, 2)])

    wb.save(out)
    print(f"\n[done] 채점 {len(rows)}문항(유효 {n}) → {out}")
    print(f"  정확도={acc:.3f} | 정답 {dist.get('정답',0)} · 부분정답 {dist.get('부분정답',0)} · 오답 {dist.get('오답',0)} · ERROR {dist.get('ERROR',0)}")
    print(f"  Macro 사실정확={mf:.2f}/2 커버리지={mc:.2f}/2 관련성={mr:.2f}/1")


if __name__ == "__main__":
    main()
