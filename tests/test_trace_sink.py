"""trace_sink 단위 테스트.

- JSONL append / xlsx append (tmp_path)
- SSE chunk text 누적
- 스트리밍 generator wrapper 가 chunk 누적 후 record 작성
- 토글 off 시 no-op (파일 생성 X)
"""

from __future__ import annotations

import asyncio
import json
import os
from typing import AsyncGenerator, List

import pytest

from app.chat.infra.rag import trace_sink
from app.chat.infra.rag.trace_sink import (
    build_trace_record,
    extract_text_from_sse_chunks,
    record_response_trace,
    wrap_stream_with_trace,
)


@pytest.fixture
def enable_trace(monkeypatch, tmp_path):
    """RESPONSE_TRACE_ENABLED=True + 디렉토리를 tmp_path 로 가리킴."""
    from app.core.config import Config

    monkeypatch.setattr(Config, "RESPONSE_TRACE_ENABLED", True, raising=False)
    monkeypatch.setattr(Config, "RESPONSE_TRACE_DIR", str(tmp_path), raising=False)
    # 새 이벤트루프마다 Lock 재초기화
    monkeypatch.setattr(trace_sink, "_XLSX_LOCK", None, raising=False)
    return tmp_path


@pytest.fixture
def disable_trace(monkeypatch, tmp_path):
    from app.core.config import Config

    monkeypatch.setattr(Config, "RESPONSE_TRACE_ENABLED", False, raising=False)
    monkeypatch.setattr(Config, "RESPONSE_TRACE_DIR", str(tmp_path), raising=False)
    return tmp_path


# ---------------------------------------------------------------------------
# SSE chunk text 추출
# ---------------------------------------------------------------------------


def _sse_chunk(content: str) -> str:
    payload = {"choices": [{"delta": {"content": content}}]}
    return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"


def test_extract_text_concatenates_delta_content():
    chunks = [_sse_chunk("안녕"), _sse_chunk("하세요"), "data: [DONE]\n\n"]
    assert extract_text_from_sse_chunks(chunks) == "안녕하세요"


def test_extract_text_ignores_done_and_empty():
    assert extract_text_from_sse_chunks([]) == ""
    assert extract_text_from_sse_chunks(["", "  ", "data: [DONE]\n\n"]) == ""


def test_extract_text_handles_message_content_fallback():
    """비스트리밍 호환 schema: choices[0].message.content"""
    payload = {"choices": [{"message": {"content": "fallback text"}}]}
    chunk = f"data: {json.dumps(payload)}\n\n"
    assert extract_text_from_sse_chunks([chunk]) == "fallback text"


def test_extract_text_skips_malformed_json():
    chunks = [
        _sse_chunk("정상"),
        "data: this is not json\n\n",
        _sse_chunk("청크"),
    ]
    result = extract_text_from_sse_chunks(chunks)
    assert "정상" in result and "청크" in result


# ---------------------------------------------------------------------------
# build_trace_record
# ---------------------------------------------------------------------------


def test_build_trace_record_has_required_fields():
    rec = build_trace_record(
        user_question="질문",
        top_docs=[{"NAME": "기초연금", "content": "내용"}],
        system_prompt="sys",
        user_message="user msg",
        intent="general",
        response_text="응답",
        is_stream=False,
        extras={"policy_priority_tag": "elderly_benefits"},
    )
    assert rec["user_question"] == "질문"
    assert rec["intent"] == "general"
    assert rec["is_stream"] is False
    assert rec["response_text"] == "응답"
    assert rec["top_docs"][0]["NAME"] == "기초연금"
    assert rec["extras"]["policy_priority_tag"] == "elderly_benefits"
    assert len(rec["trace_id"]) == 12
    assert "ts" in rec


# ---------------------------------------------------------------------------
# record_response_trace — JSONL + xlsx 생성
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_record_creates_jsonl_and_xlsx(enable_trace):
    record_response_trace(
        user_question="기초연금 알려줘",
        top_docs=[{"NAME": "기초연금"}],
        system_prompt="시스템 프롬프트",
        user_message="사용자 메시지",
        intent="general",
        response_text="답변 본문",
        extras={"detail_requested": False},
    )
    # fire-and-forget task 가 끝날 때까지 대기
    await asyncio.sleep(0.3)

    jsonl_path = os.path.join(str(enable_trace), "response_trace.jsonl")
    xlsx_path = os.path.join(str(enable_trace), "response_trace.xlsx")
    assert os.path.exists(jsonl_path), "JSONL 파일이 생성되어야 함"
    assert os.path.exists(xlsx_path), "xlsx 파일이 생성되어야 함"

    # JSONL 내용 검증
    with open(jsonl_path, "r", encoding="utf-8") as f:
        lines = [json.loads(l) for l in f.readlines() if l.strip()]
    assert len(lines) == 1
    assert lines[0]["user_question"] == "기초연금 알려줘"
    assert lines[0]["response_text"] == "답변 본문"
    assert lines[0]["is_stream"] is False

    # xlsx 내용 검증 (헤더 + 1행)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "user_question" in headers
    assert ws.max_row == 2  # 헤더 + 1 row


@pytest.mark.asyncio
async def test_record_appends_multiple_rows(enable_trace):
    for i in range(3):
        record_response_trace(
            user_question=f"질문{i}",
            top_docs=None,
            system_prompt="sys",
            user_message="msg",
            intent="general",
            response_text=f"응답{i}",
        )
    await asyncio.sleep(0.5)

    jsonl_path = os.path.join(str(enable_trace), "response_trace.jsonl")
    with open(jsonl_path, "r", encoding="utf-8") as f:
        lines = [json.loads(l) for l in f.readlines() if l.strip()]
    assert len(lines) == 3
    assert [l["user_question"] for l in lines] == ["질문0", "질문1", "질문2"]

    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(str(enable_trace), "response_trace.xlsx"))
    assert wb.active.max_row == 4  # 헤더 + 3


# ---------------------------------------------------------------------------
# 토글 off 시 no-op
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_record_no_op_when_disabled(disable_trace):
    record_response_trace(
        user_question="질문",
        top_docs=None,
        system_prompt="sys",
        user_message="msg",
        intent="general",
        response_text="응답",
    )
    await asyncio.sleep(0.1)
    assert not os.path.exists(os.path.join(str(disable_trace), "response_trace.jsonl"))
    assert not os.path.exists(os.path.join(str(disable_trace), "response_trace.xlsx"))


# ---------------------------------------------------------------------------
# wrap_stream_with_trace — chunk 누적 후 dump
# ---------------------------------------------------------------------------


async def _fake_stream(chunks: List[str]) -> AsyncGenerator[str, None]:
    for c in chunks:
        yield c


@pytest.mark.asyncio
async def test_wrap_stream_accumulates_and_dumps(enable_trace):
    chunks = [_sse_chunk("안녕"), _sse_chunk("하세요"), "data: [DONE]\n\n"]
    wrapped = wrap_stream_with_trace(
        _fake_stream(chunks),
        user_question="안부",
        top_docs=[{"NAME": "doc1"}],
        system_prompt="sys",
        user_message="msg",
        intent="general",
        extras={"detail_requested": True},
    )
    received: List[str] = []
    async for ch in wrapped:
        received.append(ch)

    # pass-through: 모든 chunk 그대로 전달
    assert received == chunks

    # dump task 가 끝날 때까지 대기
    await asyncio.sleep(0.3)

    jsonl_path = os.path.join(str(enable_trace), "response_trace.jsonl")
    with open(jsonl_path, "r", encoding="utf-8") as f:
        rec = json.loads(f.readline())
    assert rec["is_stream"] is True
    assert rec["response_text"] == "안녕하세요"
    assert rec["user_question"] == "안부"
    assert rec["extras"]["detail_requested"] is True


@pytest.mark.asyncio
async def test_wrap_stream_passthrough_when_disabled(disable_trace):
    """토글 off 시 wrap 비용 없이 pass-through, 파일 생성 X."""
    chunks = [_sse_chunk("a"), _sse_chunk("b")]
    wrapped = wrap_stream_with_trace(
        _fake_stream(chunks),
        user_question="q",
        top_docs=None,
        system_prompt="sys",
        user_message="msg",
        intent="general",
    )
    received = [ch async for ch in wrapped]
    assert received == chunks
    await asyncio.sleep(0.1)
    assert not os.path.exists(os.path.join(str(disable_trace), "response_trace.jsonl"))
