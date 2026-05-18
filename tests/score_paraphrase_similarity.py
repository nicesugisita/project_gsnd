"""Phase 2 결과 vs 원본 final_answer 유사도 비교.

Primary metric  — 사업명 Jaccard
Secondary       — LLM judge (외부 32B, JSON 출력)

입력:
  qa_test/0515_paraphrased_answers_<ts>.xlsx   (Phase 2 결과)
  qa_test/260513_gsnd_total_dataset_v0.2.xlsx  (원본 final_answer)

출력:
  qa_test/0515_similarity_<ts>.xlsx (3 sheets: per_variant, per_base, top_divergent)

사용:
  python tests/score_paraphrase_similarity.py                                    # 최신 결과 자동 선택
  python tests/score_paraphrase_similarity.py qa_test/0515_paraphrased_answers_20260517_120000.xlsx
  python tests/score_paraphrase_similarity.py qa_test/0515_paraphrased_answers_20260517_120000.xlsx 10  # 앞 10 base만 (스모크)
"""
from __future__ import annotations

import datetime as dt
import json
import re
import statistics as stat
import sys
import time
from collections import defaultdict
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ORIG_XLSX = ROOT / "qa_test" / "260513_gsnd_total_dataset_v0.2.xlsx"
OUT_DIR = ROOT / "qa_test"

LLM_URL = "http://222.122.179.212:15006/v1/chat/completions"
LLM_TIMEOUT = 120.0
LLM_RETRY_MAX = 3
LLM_RETRY_BACKOFF = 4.0

MULTITURN_PATTERN = re.compile(r"^\[(\d+)-(\d+)\]\s*")
# `- 사업명 : X` 또는 `사업명 : X` 모두 대응 (multiline).
BUSINESS_PATTERN = re.compile(r"^[\-\s\[\]\d]*\s*사업명\s*[:：]\s*(.+?)\s*$", flags=re.MULTILINE)
_NORM_NOISE = re.compile(r"[\s·.,/\-_]+")

DIVERGE_JACCARD = 0.5
DIVERGE_LLM = 70

JUDGE_SYSTEM_PROMPT = """너는 두 챗봇 응답이 같은 사용자 질문에 대해 의미적으로 동일한지 판정한다.
출력은 JSON: {"score": 0~100, "verdict": "동일|유사|부분|상이", "reason": "한 줄 사유"}.
채점 가이드:
 - 동일(90~100): 핵심 사업명 집합·결론·신청처가 모두 일치.
 - 유사(70~89): 핵심 사업명 다수 일치, 부수 사업명/순서만 다름.
 - 부분(40~69): 일부만 일치하거나, 한쪽이 더 좁은 답.
 - 상이(0~39): 다른 사업/대상/결론을 안내.
다른 텍스트 금지.
"""


# ─────────────────────────────────────────────────────────────────────
# 사업명 추출 / 정규화
# ─────────────────────────────────────────────────────────────────────

def _normalize_business(s: str) -> str:
    s = (s or "").strip()
    s = _NORM_NOISE.sub("", s)
    return s.lower()


def extract_business_set(text: str) -> set[str]:
    """응답 본문에서 '사업명' 라인을 모두 추출해 정규화 집합 반환."""
    if not text:
        return set()
    out: set[str] = set()
    for m in BUSINESS_PATTERN.finditer(text):
        name = m.group(1).strip()
        # 후행 메타 토큰 (예: "기초연금 (창원)") — 괄호 안 제거
        name = re.sub(r"\s*\(.*?\)\s*$", "", name).strip()
        if name:
            norm = _normalize_business(name)
            if norm:
                out.add(norm)
    return out


def _diff_set(a: set[str], b: set[str], orig_map: dict[str, str] | None = None) -> str:
    """차집합을 사람이 읽기 쉬운 문자열로. orig_map 으로 원문 이름 복원 가능."""
    if not a:
        return ""
    if orig_map:
        names = sorted(orig_map.get(x, x) for x in a)
    else:
        names = sorted(a)
    return " | ".join(names)


def extract_business_with_original(text: str) -> tuple[set[str], dict[str, str]]:
    """집합과 함께 (정규화 키 → 원문 이름) 매핑을 반환 — diff 출력용."""
    if not text:
        return set(), {}
    out: set[str] = set()
    orig: dict[str, str] = {}
    for m in BUSINESS_PATTERN.finditer(text):
        name = m.group(1).strip()
        name = re.sub(r"\s*\(.*?\)\s*$", "", name).strip()
        if name:
            norm = _normalize_business(name)
            if norm:
                out.add(norm)
                orig.setdefault(norm, name)
    return out, orig


# ─────────────────────────────────────────────────────────────────────
# 원본 final_answer 로딩 — (base_no, turn) → final_answer
# ─────────────────────────────────────────────────────────────────────

def load_original_answers() -> tuple[dict[tuple[int, int], str], dict[int, str]]:
    """원본 xlsx 의 모든 행을 (base_no, turn) → final_answer 로 매핑.

    추가로 base_no → original_question (turn=1) 매핑도 함께 반환.
    """
    wb = openpyxl.load_workbook(ORIG_XLSX, read_only=True, data_only=True)
    ws = wb["gsnd_total"]
    answers: dict[tuple[int, int], str] = {}
    base_q: dict[int, str] = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None:
            continue
        q_raw = str(r[1] or "").strip()
        if not q_raw:
            continue
        final = str(r[4] or "").strip() if len(r) > 4 and r[4] is not None else ""
        m = MULTITURN_PATTERN.match(q_raw)
        if m:
            base_no = int(m.group(1))
            turn = int(m.group(2))
        else:
            try:
                base_no = int(r[0])
            except (TypeError, ValueError):
                continue
            turn = 1
            base_q[base_no] = q_raw
        answers[(base_no, turn)] = final
    return answers, base_q


# ─────────────────────────────────────────────────────────────────────
# Phase 2 결과 로드
# ─────────────────────────────────────────────────────────────────────

def load_phase2_rows(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows: list[dict] = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r or r[0] is None:
            continue
        try:
            base_no = int(r[0])
            turn = int(r[1])
            variant_idx = int(r[2])
        except (TypeError, ValueError):
            continue
        rows.append({
            "base_no": base_no,
            "turn": turn,
            "variant_idx": variant_idx,
            "original_question": str(r[3] or ""),
            "paraphrased_question": str(r[4] or ""),
            "new_answer": str(r[5] or ""),
            "elapsed_sec": r[6],
            "error": str(r[7] or ""),
        })
    return rows


# ─────────────────────────────────────────────────────────────────────
# LLM judge
# ─────────────────────────────────────────────────────────────────────

def _build_judge_user(question: str, orig_ans: str, new_ans: str) -> str:
    return (
        f"[사용자 질문]\n{question}\n\n"
        f"[원본 챗봇 응답 A]\n{orig_ans or '(없음)'}\n\n"
        f"[변형 챗봇 응답 B]\n{new_ans or '(없음)'}\n"
    )


def llm_judge(client: httpx.Client, question: str, orig_ans: str, new_ans: str) -> dict:
    payload = {
        "messages": [
            {"role": "system", "content": JUDGE_SYSTEM_PROMPT},
            {"role": "user", "content": _build_judge_user(question, orig_ans, new_ans)},
        ],
        "temperature": 0,
        "max_tokens": 256,
        "stream": False,
    }
    last_err: Exception | None = None
    for attempt in range(1, LLM_RETRY_MAX + 1):
        try:
            r = client.post(LLM_URL, json=payload, timeout=LLM_TIMEOUT)
            r.raise_for_status()
            content = r.json()["choices"][0]["message"]["content"].strip()
            m = re.search(r"\{[\s\S]*\}", content)
            if not m:
                raise ValueError(f"JSON 없음: {content[:200]}")
            data = json.loads(m.group(0))
            score = data.get("score")
            verdict = str(data.get("verdict", "")).strip()
            reason = str(data.get("reason", "")).strip()
            if verdict not in ("동일", "유사", "부분", "상이"):
                raise ValueError(f"unexpected verdict: {verdict}")
            try:
                score = int(round(float(score)))
            except (TypeError, ValueError):
                raise ValueError(f"unexpected score: {score}")
            score = max(0, min(100, score))
            return {"score": score, "verdict": verdict, "reason": reason}
        except Exception as e:
            last_err = e
            if attempt < LLM_RETRY_MAX:
                wait = LLM_RETRY_BACKOFF * attempt
                print(f"    [llm retry {attempt}/{LLM_RETRY_MAX - 1}] {type(e).__name__}: {e} → {wait}s")
                time.sleep(wait)
                continue
    raise last_err  # type: ignore[misc]


# ─────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────

def find_target_xlsx() -> Path:
    if len(sys.argv) > 1 and not sys.argv[1].isdigit():
        return Path(sys.argv[1]).resolve()
    candidates = sorted(OUT_DIR.glob("0515_paraphrased_answers_*.xlsx"))
    if not candidates:
        print(f"[error] {OUT_DIR}/0515_paraphrased_answers_*.xlsx 가 없습니다.")
        sys.exit(1)
    return candidates[-1]


def main() -> None:
    target = find_target_xlsx()
    # 두번째 인자 = limit (앞에서 N개 base_no 만 처리)
    limit_idx = 2 if (len(sys.argv) > 1 and not sys.argv[1].isdigit()) else 1
    limit: int | None = None
    if len(sys.argv) > limit_idx:
        try:
            limit = int(sys.argv[limit_idx])
            if limit <= 0:
                limit = None
        except ValueError:
            limit = None

    print(f"[run] 채점 대상: {target.name}")
    if limit:
        print(f"[run] smoke limit = {limit} base_no")

    orig_answers, base_questions = load_original_answers()
    rows = load_phase2_rows(target)

    if limit:
        keep_bases = sorted({row["base_no"] for row in rows})[:limit]
        keep_set = set(keep_bases)
        rows = [r for r in rows if r["base_no"] in keep_set]

    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = OUT_DIR / f"0515_similarity_{ts}.xlsx"

    wb = openpyxl.Workbook()
    ws_pv = wb.active
    ws_pv.title = "per_variant"
    ws_pv.append([
        "base_no", "variant_idx", "turn", "paraphrased_question",
        "jaccard", "recall_orig", "missing_services", "extra_services",
        "llm_score", "llm_verdict", "llm_reason",
    ])

    per_variant_buf: list[dict] = []
    verdict_counter = {"동일": 0, "유사": 0, "부분": 0, "상이": 0}
    llm_fail = 0

    with httpx.Client() as client:
        for i, row in enumerate(rows, 1):
            base_no = row["base_no"]
            turn = row["turn"]
            variant_idx = row["variant_idx"]
            paraphrased_q = row["paraphrased_question"]
            new_ans = row["new_answer"]
            orig_ans = orig_answers.get((base_no, turn), "")

            # Jaccard
            new_set, new_orig_map = extract_business_with_original(new_ans)
            orig_set, orig_orig_map = extract_business_with_original(orig_ans)

            if not orig_set and not new_set:
                jaccard = None
                recall_orig = None
                missing = ""
                extra = ""
            else:
                inter = orig_set & new_set
                union = orig_set | new_set
                jaccard = (len(inter) / len(union)) if union else None
                recall_orig = (len(inter) / len(orig_set)) if orig_set else None
                missing = _diff_set(orig_set - new_set, set(), orig_orig_map)
                extra = _diff_set(new_set - orig_set, set(), new_orig_map)

            # LLM judge
            llm_score: int | None = None
            llm_verdict = ""
            llm_reason = ""
            if row["error"] or new_ans.startswith("[ERROR]") or new_ans.startswith("[SKIP]"):
                llm_reason = f"(skip) Phase 2 error: {row['error'] or new_ans[:80]}"
            else:
                question_for_judge = paraphrased_q or base_questions.get(base_no, "")
                try:
                    j = llm_judge(client, question_for_judge, orig_ans, new_ans)
                    llm_score = j["score"]
                    llm_verdict = j["verdict"]
                    llm_reason = j["reason"]
                    verdict_counter[llm_verdict] = verdict_counter.get(llm_verdict, 0) + 1
                except Exception as e:
                    llm_fail += 1
                    llm_reason = f"(llm fail) {type(e).__name__}: {e}"

            ws_pv.append([
                base_no, variant_idx, turn, paraphrased_q,
                round(jaccard, 4) if jaccard is not None else "N/A",
                round(recall_orig, 4) if recall_orig is not None else "N/A",
                missing, extra,
                llm_score if llm_score is not None else "",
                llm_verdict, llm_reason,
            ])
            wb.save(out_xlsx)

            per_variant_buf.append({
                "base_no": base_no, "variant_idx": variant_idx, "turn": turn,
                "paraphrased_question": paraphrased_q,
                "jaccard": jaccard, "recall_orig": recall_orig,
                "missing": missing, "extra": extra,
                "llm_score": llm_score, "llm_verdict": llm_verdict, "llm_reason": llm_reason,
                "orig_ans": orig_ans, "new_ans": new_ans,
            })

            j_str = f"{jaccard:.2f}" if jaccard is not None else "N/A"
            l_str = f"{llm_score}({llm_verdict})" if llm_score is not None else "—"
            print(f"[{i}/{len(rows)}] base={base_no} v{variant_idx} t{turn} jaccard={j_str} llm={l_str}")

    # ── per_base 집계 ────────────────────────────────────────────────
    ws_pb = wb.create_sheet("per_base")
    ws_pb.append([
        "base_no", "original_question", "n_variants",
        "jaccard_mean", "jaccard_min", "jaccard_max", "jaccard_std",
        "llm_score_mean", "llm_score_min", "llm_score_max",
        "divergence_flag",
    ])

    by_base: dict[int, list[dict]] = defaultdict(list)
    for v in per_variant_buf:
        by_base[v["base_no"]].append(v)

    per_base_summary: list[dict] = []
    for base_no in sorted(by_base):
        vs = by_base[base_no]
        jvals = [v["jaccard"] for v in vs if v["jaccard"] is not None]
        lvals = [v["llm_score"] for v in vs if v["llm_score"] is not None]

        jmean = stat.fmean(jvals) if jvals else None
        jmin = min(jvals) if jvals else None
        jmax = max(jvals) if jvals else None
        jstd = stat.pstdev(jvals) if len(jvals) > 1 else (0.0 if jvals else None)
        lmean = stat.fmean(lvals) if lvals else None
        lmin = min(lvals) if lvals else None
        lmax = max(lvals) if lvals else None

        div = False
        if jmean is not None and jmean < DIVERGE_JACCARD:
            div = True
        if lmean is not None and lmean < DIVERGE_LLM:
            div = True

        ws_pb.append([
            base_no, base_questions.get(base_no, ""), len(vs),
            round(jmean, 4) if jmean is not None else "N/A",
            round(jmin, 4) if jmin is not None else "N/A",
            round(jmax, 4) if jmax is not None else "N/A",
            round(jstd, 4) if jstd is not None else "N/A",
            round(lmean, 2) if lmean is not None else "N/A",
            lmin if lmin is not None else "N/A",
            lmax if lmax is not None else "N/A",
            "Y" if div else "",
        ])
        per_base_summary.append({
            "base_no": base_no,
            "original_question": base_questions.get(base_no, ""),
            "n_variants": len(vs),
            "jaccard_mean": jmean,
            "llm_score_mean": lmean,
            "divergence_flag": div,
            "variants": vs,
        })

    # ── top_divergent ────────────────────────────────────────────────
    ws_td = wb.create_sheet("top_divergent")
    ws_td.append([
        "base_no", "original_question", "llm_score_mean", "jaccard_mean",
        "n_variants", "missing_in_new (union)", "extra_in_new (union)",
        "orig_answer_preview", "sample_new_answer_preview",
    ])
    sortable = [b for b in per_base_summary if b["llm_score_mean"] is not None]
    sortable.sort(key=lambda x: x["llm_score_mean"])
    for b in sortable[:10]:
        miss_union: set[str] = set()
        extra_union: set[str] = set()
        sample_new = ""
        sample_orig = ""
        for v in b["variants"]:
            if v["missing"]:
                miss_union.update(v["missing"].split(" | "))
            if v["extra"]:
                extra_union.update(v["extra"].split(" | "))
            if not sample_new and v["new_ans"]:
                sample_new = v["new_ans"][:400]
            if not sample_orig and v["orig_ans"]:
                sample_orig = v["orig_ans"][:400]
        ws_td.append([
            b["base_no"], b["original_question"],
            round(b["llm_score_mean"], 2) if b["llm_score_mean"] is not None else "N/A",
            round(b["jaccard_mean"], 4) if b["jaccard_mean"] is not None else "N/A",
            b["n_variants"],
            " | ".join(sorted(miss_union)),
            " | ".join(sorted(extra_union)),
            sample_orig,
            sample_new,
        ])

    wb.save(out_xlsx)

    # ── 콘솔 리포트 ─────────────────────────────────────────────────
    all_j = [v["jaccard"] for v in per_variant_buf if v["jaccard"] is not None]
    all_r = [v["recall_orig"] for v in per_variant_buf if v["recall_orig"] is not None]
    all_l = [v["llm_score"] for v in per_variant_buf if v["llm_score"] is not None]

    print()
    print("=" * 60)
    print("[전체 변형 평균]")
    print(f"  jaccard_mean    = {stat.fmean(all_j):.4f}" if all_j else "  jaccard_mean    = N/A")
    print(f"  recall_orig_mean= {stat.fmean(all_r):.4f}" if all_r else "  recall_orig_mean= N/A")
    print(f"  llm_score_mean  = {stat.fmean(all_l):.2f}" if all_l else "  llm_score_mean  = N/A")
    print()
    print("[verdict 분포]")
    total_v = sum(verdict_counter.values()) or 1
    for k in ("동일", "유사", "부분", "상이"):
        n = verdict_counter.get(k, 0)
        print(f"  {k}: {n}  ({n * 100 / total_v:.1f}%)")
    if llm_fail:
        print(f"  (LLM 호출 실패: {llm_fail})")
    print()
    print("[발산 상위 10 base_no]")
    for b in sortable[:10]:
        lm = f"{b['llm_score_mean']:.2f}" if b['llm_score_mean'] is not None else "N/A"
        jm = f"{b['jaccard_mean']:.2f}" if b['jaccard_mean'] is not None else "N/A"
        q = (b["original_question"] or "")[:50]
        print(f"  base={b['base_no']:>3}  llm={lm}  jaccard={jm}  | {q}")
    print()
    print(f"[done] xlsx={out_xlsx}")


if __name__ == "__main__":
    main()
