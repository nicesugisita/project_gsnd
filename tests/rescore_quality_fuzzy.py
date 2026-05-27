"""품질 재채점 — 명칭차이(포함관계) 반영. 엄격(exact) vs 완화(포함관계) P/R/F1 비교.

명칭차이 인정 규칙(보수적): 두 사업명 canon이 같거나, 한쪽(길이≥4)이 다른 쪽에 통째로 포함될 때만
같은 제도로 본다. (예: '주거안정월세대출보증' ⊃ '주거안정월세대출' = 인정 / '재난적의료비' vs
'암환자의료비' = 단순 유사라 불인정.) 이렇게 해야 이름만 비슷한 다른 서비스 오인정을 막는다.

사용:
    python tests/rescore_quality_fuzzy.py <answer_v2.xlsx> <gold.xlsx> [out.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_checklist import _canon, _norm  # noqa: E402

C_Q, C_ROUND, C_DOCS = 0, 2, 14
MINLEN = 4  # 포함관계 인정 최소 길이(짧은 쪽) — 일반어 오매칭 방지

GREEN = PatternFill("solid", fgColor="C6EFCE")
HDR = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)


def _contains_match(a: str, b: str) -> bool:
    if a == b:
        return True
    m, M = sorted([a, b], key=len)
    return len(m) >= MINLEN and m in M


def _score(gen_canon, gold_canon, relaxed: bool):
    if relaxed:
        gen_hit = sum(1 for g in gen_canon if any(_contains_match(g, x) for x in gold_canon))
        gold_cov = sum(1 for x in gold_canon if any(_contains_match(g, x) for g in gen_canon))
    else:
        inter = gen_canon & gold_canon
        gen_hit = gold_cov = len(inter)
    P = gen_hit / len(gen_canon) if gen_canon else 0.0
    R = gold_cov / len(gold_canon) if gold_canon else 0.0
    F = 2 * P * R / (P + R) if (P + R) else 0.0
    return P, R, F


def main() -> None:
    if len(sys.argv) < 3:
        print("usage: python tests/rescore_quality_fuzzy.py <answer_v2.xlsx> <gold.xlsx> [out.xlsx]")
        sys.exit(1)
    ans_p, gold_p = Path(sys.argv[1]), Path(sys.argv[2])
    out = Path(sys.argv[3]) if len(sys.argv) > 3 else ans_p.with_name(ans_p.stem + "_rescore.xlsx")

    wg = openpyxl.load_workbook(gold_p, data_only=True)["gold"]
    hdr = [str(c.value or "") for c in wg[1]]
    qi = next(i for i, h in enumerate(hdr) if "question" in h)
    ai = next(i for i, h in enumerate(hdr) if "answer_doc" in h)
    gold = {}
    for r in wg.iter_rows(min_row=2, values_only=True):
        if r[qi] is None:
            continue
        names = [x.strip() for x in str(r[ai] or "").splitlines() if x.strip()]
        if names:
            gold[_norm(r[qi])] = {_canon(x) for x in names if _canon(x)}

    ws = openpyxl.load_workbook(ans_p, read_only=True, data_only=True)["stages"]
    r1 = {}  # norm(q) → R1 docs
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[C_Q] is None:
            continue
        k = _norm(r[C_Q])
        rnd = r[C_ROUND] if isinstance(r[C_ROUND], int) else 99
        if k not in r1 or rnd < r1[k][0]:
            docs = {_canon(x) for x in str(r[C_DOCS] or "").split("‖") if _canon(x.strip())}
            r1[k] = (rnd, docs, str(r[C_Q]))

    owb = Workbook(); osh = owb.active; osh.title = "재채점"
    cols = ["질문", "gold수", "검색문서수", "P(엄격)", "R(엄격)", "F1(엄격)", "P(완화)", "R(완화)", "F1(완화)", "F1Δ"]
    osh.append(cols)
    for j in range(1, len(cols) + 1):
        osh.cell(row=1, column=j).font = HDR_FONT
        osh.cell(row=1, column=j).fill = HDR
        osh.column_dimensions[get_column_letter(j)].width = 38 if j == 1 else 9

    sums = {"sP": 0, "sR": 0, "sF": 0, "rP": 0, "rR": 0, "rF": 0}
    n = 0
    improved = 0
    for k in sorted(set(r1) & set(gold)):
        gen = r1[k][1]; gc = gold[k]
        sP, sR, sF = _score(gen, gc, False)
        rP, rR, rF = _score(gen, gc, True)
        n += 1
        sums["sP"] += sP; sums["sR"] += sR; sums["sF"] += sF
        sums["rP"] += rP; sums["rR"] += rR; sums["rF"] += rF
        d = rF - sF
        if d > 0.001:
            improved += 1
        osh.append([r1[k][2], len(gc), len(gen),
                    round(sP, 2), round(sR, 2), round(sF, 2),
                    round(rP, 2), round(rR, 2), round(rF, 2), round(d, 2)])
        if d > 0.001:
            osh.cell(row=osh.max_row, column=10).fill = GREEN

    osh.freeze_panes = "B2"
    osh.auto_filter.ref = f"A1:J{osh.max_row}"
    owb.save(out)

    print(f"[done] 재채점 {n}문항 → {out}")
    print(f"{'':6}{'Precision':>11}{'Recall':>9}{'F1':>8}")
    print(f"  엄격 {sums['sP']/n:>10.3f}{sums['sR']/n:>9.3f}{sums['sF']/n:>8.3f}")
    print(f"  완화 {sums['rP']/n:>10.3f}{sums['rR']/n:>9.3f}{sums['rF']/n:>8.3f}  (명칭차이 반영)")
    print(f"  → 완화로 F1 오른 문항: {improved}/{n}")


if __name__ == "__main__":
    main()
