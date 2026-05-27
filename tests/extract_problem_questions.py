"""품질 분석 xlsx에서 '제대로 안 나온 문항'(gold 대비 F1 낮음)만 추려 정리.

입력: analyze_answer_quality.py 가 만든 *_quality.xlsx (분석 시트)
출력: 문제 문항을 F1 오름차순으로 정리한 xlsx (놓친 정답서비스·무관 노출 포함)

사용:
    python tests/extract_problem_questions.py <quality.xlsx> [out.xlsx] [F1임계(기본0.6)]
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RED = PatternFill("solid", fgColor="FFC7CE")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
HDR = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    if len(sys.argv) < 2:
        print("usage: python tests/extract_problem_questions.py <quality.xlsx> [out.xlsx] [F1임계]")
        sys.exit(1)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem.replace("_quality", "") + "_problems.xlsx")
    thr = float(sys.argv[3]) if len(sys.argv) > 3 else 0.6

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["분석"]
    H = [c.value for c in ws[1]]

    def ci(t):
        return H.index(t)

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[ci("gold有")] != "Y":
            continue
        f1 = r[ci("F1")]
        if not isinstance(f1, (int, float)):
            continue
        rows.append(r)
    problems = sorted([r for r in rows if r[ci("F1")] < thr], key=lambda r: r[ci("F1")])

    owb = Workbook()
    osh = owb.active
    osh.title = "문제문항"
    cols = [("순위", 5), ("질문", 40), ("등급", 8), ("Precision", 10), ("Recall", 9), ("F1", 8),
            ("놓친 정답서비스(미검출)", 40), ("무관 노출(정답밖)", 36),
            ("검색문서(R1)", 44), ("최종응답(R1)", 44)]
    osh.append([c[0] for c in cols])
    for j, c in enumerate(cols, 1):
        cell = osh.cell(row=1, column=j)
        cell.font = HDR_FONT; cell.fill = HDR; cell.alignment = WRAP
        osh.column_dimensions[get_column_letter(j)].width = c[1]

    for i, r in enumerate(problems, 1):
        f1 = r[ci("F1")]
        grade = "불량" if f1 < 0.3 else "주의"
        osh.append([
            i, r[ci("질문")], grade,
            r[ci("Precision")], r[ci("Recall")], r[ci("F1")],
            r[ci("빠짐(gold-only,미검출)")], r[ci("잘못(docs-only,무관후보)")],
            r[ci("검색문서(R1)")], r[ci("최종응답(R1)")],
        ])
        rr = osh.max_row
        osh.cell(row=rr, column=3).fill = RED if grade == "불량" else ORANGE
        osh.cell(row=rr, column=6).fill = RED if grade == "불량" else ORANGE
        for cc in (2, 7, 8, 9, 10):
            osh.cell(row=rr, column=cc).alignment = WRAP

    last = osh.max_row
    osh.freeze_panes = "C2"
    osh.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # 요약
    s = owb.create_sheet("요약", 0)
    s.column_dimensions["A"].width = 26; s.column_dimensions["B"].width = 12
    s.append(["품질 미흡 문항 요약"]); s["A1"].font = Font(bold=True, size=13)
    s.append(["품질 측정 문항(gold 매칭)", len(rows)])
    s.append([f"문제 문항(F1<{thr})", len(problems)])
    s.append(["  - 불량(F1<0.3)", sum(1 for r in problems if r[ci("F1")] < 0.3)])
    s.append(["  - 주의(0.3~%.1f)" % thr, sum(1 for r in problems if 0.3 <= r[ci("F1")] < thr)])
    s.append(["양호(F1>=%.1f)" % thr, len(rows) - len(problems)])
    s.append([])
    s.append(["설명", "F1<0.3=불량(거의 못 맞춤), 0.3~%.1f=주의. '놓친 정답서비스'=gold엔 있는데 안 나온 것, '무관 노출'=gold 밖인데 보여준 것." % thr])

    owb.save(out)
    print(f"[done] 문제 문항 {len(problems)}/{len(rows)} → {out}")


if __name__ == "__main__":
    main()
