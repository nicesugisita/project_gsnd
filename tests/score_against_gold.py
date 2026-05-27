"""사업명 set 골든셋 채점 — gold(기대사업명) vs 실행(consistency jsonl) → 문항별 P/R/F1.

Precision = |생성 ∩ gold| / |생성|   (이상한 사업명 안 섞였나)
Recall    = |생성 ∩ gold| / |gold|   (기대 사업명 다 나왔나)
F1        = 조화평균
사업명 매칭은 build_review_checklist._canon (연도·지역 접두어 제거 + 공백/콤마 제거) 기준.

사용:
    python tests/score_against_gold.py <gold_xlsx> <consistency_jsonl> [out_xlsx]
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_review_checklist import _service_names, _canon  # noqa: E402

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
HDR = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _split_names(cell: Any) -> List[str]:
    if cell is None:
        return []
    raw = str(cell).replace(";", "\n").replace("·", "\n").splitlines()
    return [x.strip() for x in raw if x.strip()]


def _load_gold(path: Path) -> Tuple[Dict[int, Set[str]], Dict[int, str]]:
    """no → 기대사업명 canon set (검토!=제외, 비어있지 않은 것만). 원본 표시용 map도 반환."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["gold"]
    hdr = [str(c.value or "") for c in ws[1]]

    def col(key):
        for i, h in enumerate(hdr):
            if key in h:
                return i
        return -1

    ci_no, ci_exp, ci_chk = col("no"), col("기대사업명"), col("검토")
    gold: Dict[int, Set[str]] = {}
    disp: Dict[int, str] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or ci_no < 0 or r[ci_no] is None:
            continue
        try:
            no = int(r[ci_no])
        except (TypeError, ValueError):
            continue
        chk = str(r[ci_chk]).strip() if ci_chk >= 0 and r[ci_chk] else ""
        if chk == "제외":
            continue
        names = _split_names(r[ci_exp]) if ci_exp >= 0 else []
        if not names:
            continue
        gold[no] = {_canon(n) for n in names if _canon(n)}
        disp[no] = " | ".join(names)
    return gold, disp


def _load_runs(path: Path) -> Dict[int, Dict[str, Any]]:
    """no → R1 레코드(가장 낮은 _round)."""
    best: Dict[int, Dict[str, Any]] = {}
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            no = rec.get("_gsnd_no")
            if not isinstance(no, int):
                continue
            rnd = rec.get("_round", 1)
            if no not in best or rnd < best[no].get("_round", 1):
                best[no] = rec
    return best


def _prf(gen: Set[str], gold: Set[str]) -> Tuple[float, float, float, Set[str], Set[str]]:
    inter = gen & gold
    p = len(inter) / len(gen) if gen else (1.0 if not gold else 0.0)
    r = len(inter) / len(gold) if gold else 1.0
    f1 = 2 * p * r / (p + r) if (p + r) else 0.0
    return p, r, f1, (gold - gen), (gen - gold)  # 빠짐(미검출), 잘못(무관/오답)


def main() -> None:
    if len(sys.argv) < 3:
        print("usage: python tests/score_against_gold.py <gold_xlsx> <consistency_jsonl> [out_xlsx]")
        sys.exit(1)
    gold_path, runs_path = Path(sys.argv[1]), Path(sys.argv[2])
    out = Path(sys.argv[3]) if len(sys.argv) > 3 else runs_path.with_name(runs_path.stem + "_score.xlsx")
    gold, disp = _load_gold(gold_path)
    runs = _load_runs(runs_path)
    scored = sorted(set(gold) & set(runs))
    print(f"[info] gold 라벨 {len(gold)}문항, 실행 {len(runs)}문항, 교집합(채점대상) {len(scored)}문항")
    if not scored:
        print("[error] 채점할 문항 없음 (gold 기대사업명이 채워졌는지/ no 매칭 확인)")
        sys.exit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = "score"
    cols = ["no", "질문", "gold수", "생성수", "교집합", "Precision", "Recall", "F1",
            "빠짐(미검출)", "잘못(무관/오답)", "생성 사업명"]
    ws.append(cols)
    for j, _ in enumerate(cols, 1):
        ws.cell(row=1, column=j).font = HDR_FONT
        ws.cell(row=1, column=j).fill = HDR
    widths = [6, 34, 7, 7, 7, 10, 9, 8, 26, 26, 34]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    sp = sr = sf = 0.0
    canon2disp: Dict[str, str] = {}
    for no in scored:
        rec = runs[no]
        q = rec.get("user_question", "")
        gen_names = _service_names(rec.get("response_text") or "")
        for nm in gen_names:
            canon2disp[_canon(nm)] = nm
        gen = {_canon(n) for n in gen_names if _canon(n)}
        gset = gold[no]
        p, r, f1, miss, wrong = _prf(gen, gset)
        sp += p; sr += r; sf += f1
        ws.append([
            no, q, len(gset), len(gen), len(gen & gset),
            round(p, 2), round(r, 2), round(f1, 2),
            " | ".join(sorted(miss)), " | ".join(canon2disp.get(w, w) for w in sorted(wrong)),
            " | ".join(gen_names),
        ])
        rr = ws.max_row
        for cc, val in ((6, p), (7, r), (8, f1)):
            cell = ws.cell(row=rr, column=cc)
            cell.fill = GREEN if val >= 0.8 else (YELLOW if val >= 0.5 else RED)
        for cc in (2, 9, 10, 11):
            ws.cell(row=rr, column=cc).alignment = WRAP

    nq = len(scored)
    ws.append([])
    ws.append(["[Macro 평균]", "", "", "", "", round(sp / nq, 3), round(sr / nq, 3), round(sf / nq, 3)])
    for cc in (6, 7, 8):
        ws.cell(row=ws.max_row, column=cc).font = HDR_FONT
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{nq + 1}"
    wb.save(out)
    print(f"[done] 채점 {nq}문항 → {out}")
    print(f"  Macro  Precision={sp / nq:.3f}  Recall={sr / nq:.3f}  F1={sf / nq:.3f}")


if __name__ == "__main__":
    main()
